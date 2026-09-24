import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
from pathlib import Path
from datetime import datetime, timedelta
import hashlib, os, platform, subprocess, json, base64, zipfile, shutil, tempfile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import numbers

BASE=Path(__file__).resolve().parent
DATA_BASE=Path(os.environ.get("SIGE_DATA_DIR", BASE))
DATA_BASE.mkdir(parents=True, exist_ok=True)
EXCEL=DATA_BASE/"Base_Datos_Celulares.xlsx"
ACTAS=DATA_BASE/"Actas Generadas"; ACTAS.mkdir(exist_ok=True)
TEMPLATE=DATA_BASE/"Acta de Entrega.docx"


ACTA_WORD_PS1_B64="""cGFyYW0oCiAgICBbUGFyYW1ldGVyKE1hbmRhdG9yeT0kdHJ1ZSldW3N0cmluZ10kVGVtcGxhdGUsCiAgICBbUGFyYW1ldGVyKE1hbmRhdG9yeT0kdHJ1ZSldW3N0cmluZ10kT3V0cHV0UGRmLAogICAgW1BhcmFtZXRlcihNYW5kYXRvcnk9JHRydWUpXVtzdHJpbmddJERhdGFKc29uCikKJEVycm9yQWN0aW9uUHJlZmVyZW5jZSA9ICdTdG9wJwokZGF0YSA9IEdldC1Db250ZW50IC1SYXcgLUxpdGVyYWxQYXRoICREYXRhSnNvbiB8IENvbnZlcnRGcm9tLUpzb24KJHdvcmQgPSAkbnVsbAokZG9jID0gJG51bGwKCmZ1bmN0aW9uIFJlcGxhY2UtSW5SYW5nZSgkcmFuZ2UsIFtzdHJpbmddJGZpbmRUZXh0LCBbc3RyaW5nXSRyZXBsYWNlVGV4dCwgW2Jvb2xdJGJvbGQ9JGZhbHNlKSB7CiAgICBpZiAoJG51bGwgLWVxICRyYW5nZSkgeyByZXR1cm4gJGZhbHNlIH0KICAgICRmb3VuZEFueSA9ICRmYWxzZQogICAgIyBQcmltZXJvIHNlIGludGVudGEgRmluZCBzb2JyZSBlbCByYW5nby4gRXN0byBjb25zZXJ2YSBlbCBmb3JtYXRvIGRlbCB0ZXh0byByZWVtcGxhemFkby4KICAgICR3b3JrID0gJHJhbmdlLkR1cGxpY2F0ZQogICAgJGZpbmQgPSAkd29yay5GaW5kCiAgICAkZmluZC5DbGVhckZvcm1hdHRpbmcoKQogICAgJGZpbmQuUmVwbGFjZW1lbnQuQ2xlYXJGb3JtYXR0aW5nKCkKICAgICRmaW5kLlRleHQgPSAkZmluZFRleHQKICAgICRmaW5kLkZvcndhcmQgPSAkdHJ1ZQogICAgJGZpbmQuV3JhcCA9IDAKICAgICRmaW5kLkZvcm1hdCA9ICRmYWxzZQogICAgJGZpbmQuTWF0Y2hDYXNlID0gJGZhbHNlCiAgICAkZmluZC5NYXRjaFdob2xlV29yZCA9ICRmYWxzZQogICAgJGZpbmQuTWF0Y2hXaWxkY2FyZHMgPSAkZmFsc2UKICAgIHdoaWxlICgkZmluZC5FeGVjdXRlKCkpIHsKICAgICAgICAkd29yay5UZXh0ID0gJHJlcGxhY2VUZXh0CiAgICAgICAgaWYgKCRib2xkKSB7ICR3b3JrLkZvbnQuQm9sZCA9IC0xIH0KICAgICAgICAkZm91bmRBbnkgPSAkdHJ1ZQogICAgICAgICR3b3JrLkNvbGxhcHNlKDApCiAgICAgICAgJGZpbmQgPSAkd29yay5GaW5kCiAgICAgICAgJGZpbmQuQ2xlYXJGb3JtYXR0aW5nKCkKICAgICAgICAkZmluZC5SZXBsYWNlbWVudC5DbGVhckZvcm1hdHRpbmcoKQogICAgICAgICRmaW5kLlRleHQgPSAkZmluZFRleHQKICAgICAgICAkZmluZC5Gb3J3YXJkID0gJHRydWUKICAgICAgICAkZmluZC5XcmFwID0gMAogICAgICAgICRmaW5kLkZvcm1hdCA9ICRmYWxzZQogICAgICAgICRmaW5kLk1hdGNoQ2FzZSA9ICRmYWxzZQogICAgICAgICRmaW5kLk1hdGNoV2hvbGVXb3JkID0gJGZhbHNlCiAgICAgICAgJGZpbmQuTWF0Y2hXaWxkY2FyZHMgPSAkZmFsc2UKICAgIH0KCiAgICAjIFJlc3BhbGRvIGVzcGVjaWFsbWVudGUgcGFyYSBtYXJjYWRvcmVzIGRlbnRybyBkZSBjZWxkYXM6IHNpIFdvcmQgbm8KICAgICMgZW5jdWVudHJhIGVsIG1hcmNhZG9yIHBvciBzdSBzZWdtZW50YWNpw7NuIGludGVybmEsIHJlZW1wbGF6YW1vcyBlbCB0ZXh0bwogICAgIyBjb21wbGV0byBkZSBsYSBjZWxkYS9yYW5nbyBjb25zZXJ2YW5kbyBzdXMgcMOhcnJhZm9zIHkgZWxpbWluYW5kbyBlbCBtYXJjYWRvci4KICAgIGlmICgtbm90ICRmb3VuZEFueSkgewogICAgICAgIHRyeSB7CiAgICAgICAgICAgICRyYXcgPSBbc3RyaW5nXSRyYW5nZS5UZXh0CiAgICAgICAgICAgIGlmICgkcmF3LkNvbnRhaW5zKCRmaW5kVGV4dCkpIHsKICAgICAgICAgICAgICAgICRuZXdUZXh0ID0gJHJhdy5SZXBsYWNlKCRmaW5kVGV4dCwgJHJlcGxhY2VUZXh0KQogICAgICAgICAgICAgICAgJHJhbmdlLlRleHQgPSAkbmV3VGV4dAogICAgICAgICAgICAgICAgaWYgKCRib2xkKSB7ICRyYW5nZS5Gb250LkJvbGQgPSAtMSB9CiAgICAgICAgICAgICAgICAkZm91bmRBbnkgPSAkdHJ1ZQogICAgICAgICAgICB9CiAgICAgICAgfSBjYXRjaCB7fQogICAgfQogICAgcmV0dXJuICRmb3VuZEFueQp9CnRyeSB7CiAgICAkd29yZCA9IE5ldy1PYmplY3QgLUNvbU9iamVjdCBXb3JkLkFwcGxpY2F0aW9uCiAgICAkd29yZC5WaXNpYmxlID0gJGZhbHNlCiAgICAkd29yZC5EaXNwbGF5QWxlcnRzID0gMAogICAgIyBMYSBwbGFudGlsbGEgb3JpZ2luYWwgc2UgYWJyZSBzb2xvIHBhcmEgbGVjdHVyYSB5IG51bmNhIHNlIGd1YXJkYS4KICAgICRkb2MgPSAkd29yZC5Eb2N1bWVudHMuT3BlbigkVGVtcGxhdGUsICRmYWxzZSwgJHRydWUpCgogICAgZm9yZWFjaCAoJHByb3AgaW4gJGRhdGEuUFNPYmplY3QuUHJvcGVydGllcykgewogICAgICAgICRmaW5kVGV4dCA9IFtzdHJpbmddJHByb3AuTmFtZQogICAgICAgICRyZXBsYWNlVGV4dCA9IFtzdHJpbmddJHByb3AuVmFsdWUKICAgICAgICAkaXNEYXRlID0gKCRmaW5kVGV4dCAtZXEgJ3t7RmVjaGEgZGUgZW50cmVnYX19JykKCiAgICAgICAgIyBDdWVycG8gcHJpbmNpcGFsIGRlbCBkb2N1bWVudG8uCiAgICAgICAgUmVwbGFjZS1JblJhbmdlICRkb2MuQ29udGVudCAkZmluZFRleHQgJHJlcGxhY2VUZXh0ICRpc0RhdGUKCiAgICAgICAgIyBUYWJsYXM6IGxvcyBjYW1wb3MgZGVsIGFjdGEgZXN0w6FuIGRlbnRybyBkZSBjZWxkYXMgZGUgdGFibGEuCiAgICAgICAgZm9yZWFjaCAoJHRhYmxlIGluICRkb2MuVGFibGVzKSB7CiAgICAgICAgICAgIGZvcmVhY2ggKCRyb3cgaW4gJHRhYmxlLlJvd3MpIHsKICAgICAgICAgICAgICAgIGZvcmVhY2ggKCRjZWxsIGluICRyb3cuQ2VsbHMpIHsKICAgICAgICAgICAgICAgICAgICAkY2VsbFJhbmdlID0gJGNlbGwuUmFuZ2UuRHVwbGljYXRlCiAgICAgICAgICAgICAgICAgICAgUmVwbGFjZS1JblJhbmdlICRjZWxsUmFuZ2UgJGZpbmRUZXh0ICRyZXBsYWNlVGV4dCAkaXNEYXRlCiAgICAgICAgICAgICAgICB9CiAgICAgICAgICAgIH0KICAgICAgICB9CgogICAgICAgICMgRW5jYWJlemFkb3MsIHBpZXMgeSBvdHJvcyBTdG9yeVJhbmdlcy4KICAgICAgICBmb3JlYWNoICgkc3RvcnkgaW4gJGRvYy5TdG9yeVJhbmdlcykgewogICAgICAgICAgICAkcmFuZ2UgPSAkc3RvcnkuRHVwbGljYXRlCiAgICAgICAgICAgIHdoaWxlICgkbnVsbCAtbmUgJHJhbmdlKSB7CiAgICAgICAgICAgICAgICBSZXBsYWNlLUluUmFuZ2UgJHJhbmdlICRmaW5kVGV4dCAkcmVwbGFjZVRleHQgJGlzRGF0ZQogICAgICAgICAgICAgICAgdHJ5IHsgJHJhbmdlID0gJHJhbmdlLk5leHRTdG9yeVJhbmdlIH0gY2F0Y2ggeyAkcmFuZ2UgPSAkbnVsbCB9CiAgICAgICAgICAgIH0KICAgICAgICB9CgogICAgICAgICMgQ3VhZHJvcyBkZSB0ZXh0by9TaGFwZXMuCiAgICAgICAgZm9yZWFjaCAoJHNoYXBlIGluICRkb2MuU2hhcGVzKSB7CiAgICAgICAgICAgIHRyeSB7CiAgICAgICAgICAgICAgICBpZiAoJHNoYXBlLlRleHRGcmFtZS5IYXNUZXh0KSB7CiAgICAgICAgICAgICAgICAgICAgUmVwbGFjZS1JblJhbmdlICRzaGFwZS5UZXh0RnJhbWUuVGV4dFJhbmdlICRmaW5kVGV4dCAkcmVwbGFjZVRleHQgJGlzRGF0ZQogICAgICAgICAgICAgICAgfQogICAgICAgICAgICB9IGNhdGNoIHt9CiAgICAgICAgfQogICAgfQoKICAgICMgVmVyaWZpY2FjacOzbiBlc3RyaWN0YSBhbnRlcyBkZSBjcmVhciBlbCBQREY6IG5pbmfDum4gbWFyY2Fkb3IgcHVlZGUgcXVlZGFyCiAgICAjIGVuIGN1ZXJwbywgdGFibGFzLCBlbmNhYmV6YWRvcy9waWVzIG8gY3VhZHJvcyBkZSB0ZXh0by4KICAgIGZvcmVhY2ggKCRwcm9wIGluICRkYXRhLlBTT2JqZWN0LlByb3BlcnRpZXMpIHsKICAgICAgICAkbGVmdCA9IFtzdHJpbmddJHByb3AuTmFtZQogICAgICAgICRyZW1haW5pbmcgPSAkZmFsc2UKICAgICAgICB0cnkgewogICAgICAgICAgICAkY2hlY2sgPSAkZG9jLkNvbnRlbnQuRHVwbGljYXRlCiAgICAgICAgICAgICRmID0gJGNoZWNrLkZpbmQKICAgICAgICAgICAgJGYuQ2xlYXJGb3JtYXR0aW5nKCk7ICRmLlRleHQ9JGxlZnQ7ICRmLkZvcndhcmQ9JHRydWU7ICRmLldyYXA9MDsgJGYuTWF0Y2hXaWxkY2FyZHM9JGZhbHNlCiAgICAgICAgICAgIGlmICgkZi5FeGVjdXRlKCkpIHsgJHJlbWFpbmluZyA9ICR0cnVlIH0KICAgICAgICB9IGNhdGNoIHt9CiAgICAgICAgaWYgKC1ub3QgJHJlbWFpbmluZykgewogICAgICAgICAgICBmb3JlYWNoICgkdGFibGUgaW4gJGRvYy5UYWJsZXMpIHsKICAgICAgICAgICAgICAgIGZvcmVhY2ggKCRyb3cgaW4gJHRhYmxlLlJvd3MpIHsKICAgICAgICAgICAgICAgICAgICBmb3JlYWNoICgkY2VsbCBpbiAkcm93LkNlbGxzKSB7CiAgICAgICAgICAgICAgICAgICAgICAgIGlmIChbc3RyaW5nXSRjZWxsLlJhbmdlLlRleHQgLWxpa2UgIiokbGVmdCoiKSB7ICRyZW1haW5pbmcgPSAkdHJ1ZTsgYnJlYWsgfQogICAgICAgICAgICAgICAgICAgIH0KICAgICAgICAgICAgICAgICAgICBpZiAoJHJlbWFpbmluZykgeyBicmVhayB9CiAgICAgICAgICAgICAgICB9CiAgICAgICAgICAgICAgICBpZiAoJHJlbWFpbmluZykgeyBicmVhayB9CiAgICAgICAgICAgIH0KICAgICAgICB9CiAgICAgICAgaWYgKCRyZW1haW5pbmcpIHsgdGhyb3cgIkVsIG1hcmNhZG9yICRsZWZ0IG5vIGZ1ZSByZWVtcGxhemFkbyBlbiBsYSBwbGFudGlsbGEgZGUgV29yZC4iIH0KICAgIH0KCiAgICAkZG9jLkV4cG9ydEFzRml4ZWRGb3JtYXQoJE91dHB1dFBkZiwgMTcsICRmYWxzZSwgMCwgMCwgMCwgMCwgMCwgJHRydWUsICRmYWxzZSwgMCwgJHRydWUsICR0cnVlLCAkZmFsc2UpCn0KZmluYWxseSB7CiAgICBpZiAoJGRvYyAtbmUgJG51bGwpIHsKICAgICAgICB0cnkgeyAkZG9jLkNsb3NlKCRmYWxzZSkgfSBjYXRjaCB7fQogICAgICAgIFt2b2lkXVtTeXN0ZW0uUnVudGltZS5JbnRlcm9wU2VydmljZXMuTWFyc2hhbF06OlJlbGVhc2VDb21PYmplY3QoJGRvYykKICAgIH0KICAgIGlmICgkd29yZCAtbmUgJG51bGwpIHsKICAgICAgICB0cnkgeyAkd29yZC5RdWl0KCkgfSBjYXRjaCB7fQogICAgICAgIFt2b2lkXVtTeXN0ZW0uUnVudGltZS5JbnRlcm9wU2VydmljZXMuTWFyc2hhbF06OlJlbGVhc2VDb21PYmplY3QoJHdvcmQpCiAgICB9CiAgICBbR0NdOjpDb2xsZWN0KCk7IFtHQ106OldhaXRGb3JQZW5kaW5nRmluYWxpemVycygpCn0K"""

def materialize_acta_resources(ps_path):
    """Crea temporalmente solo el script PS1 desde el contenido integrado."""
    ps_path.write_bytes(base64.b64decode(ACTA_WORD_PS1_B64))

HEADERS=["ID","Operador","Razon Social","Nombre y Apellido","Correo","Cargo","DNI","Número celular","Plan Celular","Tipo de Plan","Modelo Asignado","Gama Celular","IMEI","Estado","Estado de equipo","Fecha Asignacion","Fecha Liberacion","Observaciones"]
STATES=["Disponible","Asignado","Mantenimiento","Baja"]
EQUIP_STATES=["Equipo nuevo","Equipo seminuevo"]
OPERADORES=["Claro","Movistar","Entel","Bitel"]
RAZONES=["Inchcape Motors Perú S.A.","Inchcape Automotriz Perú S.A.","Derco Perú S.A.","Dercocenter S.A.C"]
PLANES=["31.90","44.90","52.90","55.90","59.90","62.90","79.90","98.90","99.90","105.90","109.90","119.90","129.90","199.90","201.50","249.00","NA"]
TIPOS_PLAN=["Plan 1","Plan 2 y 3","Plan 4","BAM","Chip","NA"]
GAMAS=["Gama 1","Gama 2","Gama 3","Gama 4","Gama 5","NA"]
ROLES=["Administrador","Usuario Asignador","Consulta"]

ASSIGN_HEADERS=["Operador","Razón Social","Nombre y Apellidos","Correo","Cargo","DNI","Número celular","Plan Celular","Tipo de Plan","Marca y Modelo","Tipo de Gama","IMEI","Situación"]
EQUIP_HEADERS=["IMEI","Marca y Modelo","Gama Celular","Estado de equipo","Situación","Fecha alta","Observaciones"]
ACTAS_ENTREGA_HEADERS=["Fecha","Usuario","Nombre","DNI","Fecha de Entrega","Nombre del documento generado"]

def sha(s): return hashlib.sha256(s.encode()).hexdigest()
def now(): return datetime.now().strftime("%d/%m/%Y")
def book(): return load_workbook(EXCEL)

def ensure_equipment_schema():
    """Migra Inventario Equipos: elimina Marca/Modelo y usa Marca y Modelo."""
    if not EXCEL.exists(): return
    wb=load_workbook(EXCEL)
    if "Inventario Equipos" not in wb.sheetnames: return
    ws=wb["Inventario Equipos"]
    target=["ID Equipo","IMEI","Marca y Modelo","Gama Celular","Estado de equipo","Situación","Fecha alta","Observaciones"]
    current=[str(c.value or "").strip() for c in ws[1]]
    if current!=target:
        old_rows=[list(r) for r in ws.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
        old=current
        ws.delete_rows(1,ws.max_row)
        ws.append(target)
        for r in old_rows:
            m={h:(r[i] if i<len(r) else None) for i,h in enumerate(old)}
            combined=m.get("Marca y Modelo")
            if not combined:
                marca=str(m.get("Marca") or "").strip()
                modelo=str(m.get("Modelo") or "").strip()
                combined=" ".join(x for x in (marca,modelo) if x)
            ws.append([m.get("ID Equipo"),m.get("IMEI"),combined,m.get("Gama Celular"),m.get("Estado de equipo"),m.get("Situación"),m.get("Fecha alta"),m.get("Observaciones")])
    # Sincroniza el nombre del equipo combinado hacia las asignaciones e inventario.
    equip_by_id={}
    for r in ws.iter_rows(min_row=2):
        if r[0].value not in (None, ""):
            equip_by_id[str(r[0].value).strip()]=str(r[2].value or "").strip()
    if "Inventario" in wb.sheetnames:
        wi=wb["Inventario"]
        for r in wi.iter_rows(min_row=2):
            eqid=str(r[0].value or "").strip()
            if eqid in equip_by_id and equip_by_id[eqid]: r[10].value=equip_by_id[eqid]
    if "Asignaciones" in wb.sheetnames:
        wa=wb["Asignaciones"]
        for r in wa.iter_rows(min_row=2):
            eqid=str(r[1].value or "").strip()
            if eqid in equip_by_id and equip_by_id[eqid]: r[6].value=equip_by_id[eqid]
    wb.save(EXCEL)

def ensure_assignments_schema():
    """Asegura que Asignaciones contenga todos los datos necesarios para el acta."""
    if not EXCEL.exists():
        return
    wb=load_workbook(EXCEL)
    if "Asignaciones" not in wb.sheetnames:
        wb.create_sheet("Asignaciones")
    wa=wb["Asignaciones"]
    target=["ID Asignación","ID Equipo","Nombre y Apellido","DNI","Correo","Cargo","Marca y Modelo","IMEI","Estado de equipo","Número celular","Operador","Fecha Asignación","Fecha Liberación","Estado","Detalle"]
    current=[str(c.value or "").strip() for c in wa[1]]
    if current!=target:
        old_rows=[]
        if any(current):
            old_rows=[list(r) for r in wa.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
        wa.delete_rows(1,wa.max_row)
        wa.append(target)
        for r in old_rows:
            m={h:(r[i] if i<len(r) else None) for i,h in enumerate(current)}
            wa.append([m.get("ID Asignación"),m.get("ID Equipo"),m.get("Nombre y Apellido"),m.get("DNI"),m.get("Correo"),m.get("Cargo"),m.get("Marca y Modelo") or m.get("Modelo Asignado"),m.get("IMEI"),m.get("Estado de equipo"),m.get("Número celular"),m.get("Operador"),m.get("Fecha Asignación"),m.get("Fecha Liberación"),m.get("Estado"),m.get("Detalle")])
    # Completar columnas nuevas de asignaciones antiguas usando el ID Equipo.
    inv_by_id={}
    if "Inventario" in wb.sheetnames:
        for r in wb["Inventario"].iter_rows(min_row=2):
            key=str(r[0].value or "").strip()
            if key: inv_by_id[key]=r
    for ar in wa.iter_rows(min_row=2):
        eq=str(ar[1].value or "").strip()
        inv=inv_by_id.get(eq)
        if inv is not None:
            vals={5:inv[5].value,6:inv[10].value,7:inv[12].value,8:inv[14].value,9:inv[7].value,10:inv[1].value}
            for col,val in vals.items():
                if ar[col].value in (None,"") and val not in (None,""): ar[col].value=val
    wb.save(EXCEL)

def ensure_movimientos_schema():
    wb=load_workbook(EXCEL)
    ws=wb["Movimientos"]
    headers=[str(c.value or "").strip() for c in ws[1]]
    target=["ID Movimiento","ID Equipo","Tipo Movimiento","Fecha","Usuario","IMEI","Detalle"]
    if headers == target:
        return
    # Migración V20.2 -> V20.3: conserva los datos existentes y agrega IMEI.
    old_rows=list(ws.iter_rows(min_row=2,values_only=True))
    old_headers=headers
    ws.delete_rows(1,ws.max_row)
    ws.append(target)
    for row in old_rows:
        d=dict(zip(old_headers,row))
        eq=d.get("ID Equipo","")
        imei=d.get("IMEI","") or ""
        if not imei and eq not in (None, ""):
            eqs=str(eq).strip()
            for sheet_name,id_col,imei_col in (("Inventario Equipos",0,1),("Inventario",0,12)):
                if sheet_name in wb.sheetnames:
                    for rr in wb[sheet_name].iter_rows(min_row=2):
                        if str(rr[id_col].value or "").strip()==eqs:
                            imei=str(rr[imei_col].value or "").strip()
                            if imei: break
                    if imei: break
        ws.append([d.get("ID Movimiento",""),eq,d.get("Tipo Movimiento",""),d.get("Fecha",""),d.get("Usuario",""),imei,d.get("Detalle","")])
    wb.save(EXCEL)

def ensure_actas_entrega_schema():
    """Crea y mantiene la hoja Actas de Entrega para la trazabilidad y apertura de PDFs."""
    if not EXCEL.exists():
        return
    wb=load_workbook(EXCEL)
    if "Actas de Entrega" not in wb.sheetnames:
        ws=wb.create_sheet("Actas de Entrega")
        ws.append(ACTAS_ENTREGA_HEADERS)
        # Migra el historial existente de Actas cuando sea posible.
        if "Actas" in wb.sheetnames:
            old=wb["Actas"]
            for r in old.iter_rows(min_row=2,values_only=True):
                if not any(v not in (None,"") for v in r):
                    continue
                # La hoja histórica V20 no almacenaba el usuario generador.
                # No se inventa el dato: se deja explícitamente como no registrado.
                fecha=str(r[5] or "") if len(r)>5 else ""
                nombre=str(r[3] or "") if len(r)>3 else ""
                dni=str(r[4] or "") if len(r)>4 else ""
                archivo=str(r[6] or "") if len(r)>6 else ""
                ws.append([fecha,"No registrado",nombre,dni,fecha,archivo])
    else:
        ws=wb["Actas de Entrega"]
        current=[str(c.value or "").strip() for c in ws[1]] if ws.max_row else []
        if current!=ACTAS_ENTREGA_HEADERS:
            rows=[list(r) for r in ws.iter_rows(min_row=2,values_only=True) if any(v not in (None,"") for v in r)]
            ws.delete_rows(1,ws.max_row)
            ws.append(ACTAS_ENTREGA_HEADERS)
            for r in rows:
                m={h:(r[i] if i<len(r) else None) for i,h in enumerate(current)}
                ws.append([m.get("Fecha"),m.get("Usuario","No registrado") or "No registrado",m.get("Nombre"),m.get("DNI"),m.get("Fecha de Entrega") or m.get("Fecha"),m.get("Nombre del documento generado") or m.get("Archivo PDF")])
    wb.save(EXCEL)

def format_database_headers():
    """Aplica formato profesional a los títulos de todas las hojas de la base."""
    if not EXCEL.exists():
        return
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb=load_workbook(EXCEL)
    fill=PatternFill("solid", fgColor="1F4E78")
    font=Font(bold=True, color="FFFFFF")
    side=Side(style="thin", color="D9E2F3")
    border=Border(bottom=side)
    for ws in wb.worksheets:
        if ws.max_column < 1:
            continue
        for cell in ws[1]:
            if cell.value is not None:
                cell.font=font
                cell.fill=fill
                cell.alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border=border
        ws.row_dimensions[1].height=30
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions if ws.max_row >= 1 else None
        for col in ws.columns:
            letter=col[0].column_letter
            maxlen=max((len(str(c.value or "")) for c in col),default=0)
            ws.column_dimensions[letter].width=min(max(maxlen+2,12),35)
    wb.save(EXCEL)

def normalize_application_dates():
    """Normaliza fechas existentes y nuevas al formato dd/mm/aaaa, sin hora."""
    if not EXCEL.exists(): return
    wb=load_workbook(EXCEL)
    mappings={
        "Inventario":[15,16],
        "Inventario Equipos":[6],
        "Asignaciones":[11,12],
        "Movimientos":[3],
        "Actas":[5],
        "Actas de Entrega":[0,4],
    }
    changed=False
    for sheet, indexes in mappings.items():
        if sheet not in wb.sheetnames: continue
        ws=wb[sheet]
        for row in ws.iter_rows(min_row=2):
            for idx in indexes:
                cell=row[idx]
                if cell.value not in (None, ""):
                    value=fmt_date(cell.value)
                    if value and str(cell.value) != value:
                        cell.value=value
                        changed=True
    if changed: wb.save(EXCEL)

def log(user,eq,typ,detail,imei=None):
    wb=book(); ws=wb["Movimientos"]
    # V20.3: la hoja Movimientos incorpora IMEI entre Usuario y Detalle.
    # Si el llamador no lo entrega, se intenta resolver por ID Equipo o por IMEI.
    if not imei:
        eqs=str(eq or "").strip()
        if eqs:
            for sheet_name, id_col, imei_col in (("Inventario Equipos",0,1),("Inventario",0,12)):
                if sheet_name in wb.sheetnames:
                    for rr in wb[sheet_name].iter_rows(min_row=2):
                        if str(rr[id_col].value or "").strip()==eqs:
                            imei=str(rr[imei_col].value or "").strip()
                            if imei:
                                break
                    if imei:
                        break
        # Rehabilitación históricamente envía el IMEI como segundo argumento.
        if not imei and eqs.isdigit() and len(eqs) in (14,15,16,17,18,19,20):
            imei=eqs
    n=max([int(r[0].value or 0) for r in ws.iter_rows(min_row=2)],default=0)+1
    ws.append([n,eq,typ,now(),user,imei or "",detail]); wb.save(EXCEL)

def open_file(p):
    try:
        if platform.system()=="Windows": os.startfile(str(p))
        elif platform.system()=="Darwin": subprocess.Popen(["open",str(p)])
        else: subprocess.Popen(["xdg-open",str(p)])
    except Exception: pass

def fmt_date(v):
    if not v: return ""
    if isinstance(v, datetime): return v.strftime("%d/%m/%Y")
    text=str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try: return datetime.strptime(text,fmt).strftime("%d/%m/%Y")
        except Exception: pass
    try: return datetime.fromisoformat(text).strftime("%d/%m/%Y")
    except Exception: return text.split(" ")[0]

class Login(tk.Tk):
    def __init__(self):
        ensure_assignments_schema()
        ensure_equipment_schema()
        ensure_actas_entrega_schema()
        ensure_movimientos_schema()
        normalize_application_dates()
        super().__init__(); self.title("Gestión Corporativa de Celulares"); self.geometry("430x280"); self.resizable(False,False)
        ttk.Label(self,text="GESTIÓN CORPORATIVA DE CELULARES",font=("Segoe UI",14,"bold")).pack(pady=28)
        f=ttk.Frame(self); f.pack()
        ttk.Label(f,text="Usuario").grid(row=0,column=0,padx=8,pady=8); self.u=ttk.Entry(f); self.u.grid(row=0,column=1)
        ttk.Label(f,text="Clave").grid(row=1,column=0,padx=8,pady=8); self.p=ttk.Entry(f,show="*"); self.p.grid(row=1,column=1)
        ttk.Button(self,text="Ingresar",command=self.login).pack(pady=18)
        self.bind("<Return>",lambda e:self.login())
    def login(self):
        wb=book()
        for r in wb["Usuarios"].iter_rows(min_row=2,values_only=True):
            if r[0]==self.u.get().strip() and r[2]==sha(self.p.get()):
                self.destroy(); App(r[0],r[1]).mainloop(); return
        messagebox.showerror("Acceso","Usuario o clave incorrectos.")

class App(tk.Tk):
    def __init__(self,user,role):
        super().__init__(); self.user=user; self.role=role
        self.title(f"Gestión Corporativa de Celulares - {role}"); self.geometry("1400x760")
        top=ttk.Frame(self); top.pack(fill="x",padx=12,pady=8)
        ttk.Label(top,text=f"Usuario: {user} | Perfil: {role}",font=("Segoe UI",10,"bold")).pack(side="left")
        ttk.Button(top,text="Cerrar sesión",command=self.logout).pack(side="right")
        self.nb=ttk.Notebook(self); self.nb.pack(fill="both",expand=True,padx=10,pady=5)
        self.build_dashboard(); self.build_inventory(); self.build_equipment_inventory(); self.build_history(); self.build_actas_entrega()
        if role=="Administrador": self.build_users()
        self.refresh_all()
    def level(self): return {"Consulta":1,"Usuario Asignador":2,"Administrador":3}[self.role]
    def logout(self): self.destroy(); Login().mainloop()

    def build_dashboard(self):
        t=ttk.Frame(self.nb); self.nb.add(t,text="Dashboard"); self.cards={}
        cards=ttk.Frame(t); cards.pack(fill="x",padx=8,pady=(12,6))
        for i,s in enumerate(["Total equipos","Disponible","Asignado","Mantenimiento","Baja"]):
            b=ttk.LabelFrame(cards,text=s); b.grid(row=0,column=i,padx=8,pady=6,sticky="nsew")
            self.cards[s]=ttk.Label(b,text="0",font=("Segoe UI",22,"bold")); self.cards[s].pack(padx=28,pady=18)
            cards.columnconfigure(i,weight=1)

        charts=ttk.Frame(t); charts.pack(fill="both",expand=True,padx=10,pady=(4,10))
        daily_box=ttk.LabelFrame(charts,text="Movimiento diario de equipos - últimos 14 días")
        daily_box.grid(row=0,column=0,padx=(0,6),pady=4,sticky="nsew")
        weekly_box=ttk.LabelFrame(charts,text="Movimiento semanal de equipos - últimas 8 semanas")
        weekly_box.grid(row=0,column=1,padx=(6,0),pady=4,sticky="nsew")
        charts.columnconfigure(0,weight=1); charts.columnconfigure(1,weight=1); charts.rowconfigure(0,weight=1)
        self.daily_chart=tk.Canvas(daily_box,height=330,bg="white",highlightthickness=0)
        self.daily_chart.pack(fill="both",expand=True,padx=8,pady=8)
        self.weekly_chart=tk.Canvas(weekly_box,height=330,bg="white",highlightthickness=0)
        self.weekly_chart.pack(fill="both",expand=True,padx=8,pady=8)
        self.daily_chart.bind("<Configure>",lambda e:self.refresh_movement_charts())
        self.weekly_chart.bind("<Configure>",lambda e:self.refresh_movement_charts())

    def build_inventory(self):
        self.tabi=ttk.Frame(self.nb); self.nb.add(self.tabi,text="Asignaciones")
        bar=ttk.Frame(self.tabi); bar.pack(fill="x",pady=6)
        ttk.Label(bar,text="Buscar").pack(side="left"); self.search=ttk.Entry(bar,width=30); self.search.pack(side="left",padx=6); self.search.bind("<KeyRelease>",lambda e:self.refresh_inventory())
        buttons=[("Nueva asignación",self.assign),("Recuperar",self.release),("Mantenimiento",self.maintenance),("Baja",self.baja),("Generar documento",self.acta),("Exportar XLSX",self.export_xlsx)]
        if self.level()>=2:
            for text,cmd in buttons: ttk.Button(bar,text=text,command=cmd).pack(side="left",padx=2)
        if self.level()==3: ttk.Button(bar,text="Eliminar",command=self.delete).pack(side="left",padx=2)
        self.tree=ttk.Treeview(self.tabi,columns=ASSIGN_HEADERS,show="headings")
        for c in ASSIGN_HEADERS: self.tree.heading(c,text=c); self.tree.column(c,width=110)
        self.tree.pack(fill="both",expand=True)

    def build_equipment_inventory(self):
        t=ttk.Frame(self.nb); self.nbe=t; self.nb.add(t,text="Inventario Equipos / IMEI")
        bar=ttk.Frame(t); bar.pack(fill="x",pady=6)
        if self.level()>=2:
            ttk.Button(bar,text="Registrar equipo / IMEI",command=self.equipment_form).pack(side="left",padx=2)
            ttk.Button(bar,text="Importar Excel",command=self.bulk).pack(side="left",padx=2)
            ttk.Button(bar,text="Rehabilitar equipo",command=self.rehabilitate).pack(side="left",padx=2)
        ttk.Button(bar,text="¿Cómo importar?",command=self.import_help).pack(side="left",padx=2)
        self.et=ttk.Treeview(t,columns=EQUIP_HEADERS,show="headings")
        for c in EQUIP_HEADERS:self.et.heading(c,text=c);self.et.column(c,width=135)
        self.et.pack(fill="both",expand=True)

    def build_history(self):
        t=ttk.Frame(self.nb); self.nb.add(t,text="Historial")
        cols=["Tipo Movimiento","Fecha","Usuario","IMEI","Detalle"]; self.ht=ttk.Treeview(t,columns=cols,show="headings")
        widths={"Tipo Movimiento":170,"Fecha":110,"Usuario":150,"IMEI":170,"Detalle":420}
        for c in cols:self.ht.heading(c,text=c);self.ht.column(c,width=widths.get(c,170))
        self.ht.pack(fill="both",expand=True)

    def build_actas_entrega(self):
        t=ttk.Frame(self.nb); self.nba=t; self.nb.add(t,text="Actas de Entrega")
        bar=ttk.Frame(t); bar.pack(fill="x",pady=6)
        ttk.Label(bar,text="Doble clic sobre un acta para abrir el PDF",font=("Segoe UI",9,"italic")).pack(side="left",padx=4)
        ttk.Button(bar,text="Actualizar",command=self.refresh_actas_entrega).pack(side="right",padx=2)
        self.at=ttk.Treeview(t,columns=ACTAS_ENTREGA_HEADERS,show="headings")
        widths={"Fecha":110,"Usuario":150,"Nombre":220,"DNI":110,"Fecha de Entrega":130,"Nombre del documento generado":360}
        for c in ACTAS_ENTREGA_HEADERS:
            self.at.heading(c,text=c)
            self.at.column(c,width=widths.get(c,160),anchor="w")
        self.at.pack(fill="both",expand=True)
        self.at.bind("<Double-1>",self.open_selected_acta)

    def build_users(self):
        t=ttk.Frame(self.nb); self.nb.add(t,text="Usuarios")
        f=ttk.Frame(t);f.pack(fill="x",pady=8)
        self.eu=ttk.Entry(f);self.ep=tk.Entry(f,show="*");self.er=ttk.Combobox(f,values=ROLES,state="readonly");self.er.set("Consulta")
        for lab,w in [("Usuario",self.eu),("Clave",self.ep),("Rol",self.er)]: ttk.Label(f,text=lab).pack(side="left",padx=4);w.pack(side="left",padx=4)
        ttk.Button(f,text="Crear usuario",command=self.create_user).pack(side="left",padx=5);ttk.Button(f,text="Eliminar usuario",command=self.delete_user).pack(side="left")
        self.ut=ttk.Treeview(t,columns=("Usuario","Rol"),show="headings")
        for c in ("Usuario","Rol"):self.ut.heading(c,text=c)
        self.ut.pack(fill="both",expand=True)

    def rows(self): return [list(r) for r in book()["Inventario"].iter_rows(min_row=2,values_only=True)]
    def equipment_rows(self): return [list(r) for r in book()["Inventario Equipos"].iter_rows(min_row=2,values_only=True)]
    def refresh_all(self): self.refresh_inventory();self.refresh_equipment();self.refresh_dashboard();self.refresh_history();self.refresh_actas_entrega();self.refresh_users()

    def refresh_dashboard(self):
        rows=self.equipment_rows()
        self.cards["Total equipos"].config(text=str(len(rows)))
        for x in STATES:self.cards[x].config(text=str(sum(1 for r in rows if str(r[5])==x)))
        self.refresh_movement_charts()

    def _movement_dates(self):
        dates=[]
        wb=book()
        if "Movimientos" not in wb.sheetnames: return dates
        for r in wb["Movimientos"].iter_rows(min_row=2,values_only=True):
            value=r[3] if len(r)>3 else None
            if isinstance(value,datetime): dates.append(value.date()); continue
            if value in (None,""): continue
            txt=str(value).strip().split()[0]
            parsed=None
            for fmt in ("%d/%m/%Y","%Y-%m-%d","%d-%m-%Y"):
                try: parsed=datetime.strptime(txt,fmt).date(); break
                except ValueError: pass
            if parsed: dates.append(parsed)
        return dates

    def _draw_bar_chart(self,canvas,labels,values):
        canvas.delete("all")
        w=max(canvas.winfo_width(),420); h=max(canvas.winfo_height(),260)
        left,top,right,bottom=48,22,18,55
        pw=max(w-left-right,10); ph=max(h-top-bottom,10)
        maxv=max(values+[1])
        canvas.create_line(left,top,left,h-bottom,fill="#777")
        canvas.create_line(left,h-bottom,w-right,h-bottom,fill="#777")
        for tick in range(5):
            y=h-bottom-(ph*tick/4)
            val=round(maxv*tick/4)
            canvas.create_line(left-4,y,w-right,y,fill="#e7e7e7")
            canvas.create_text(left-8,y,text=str(val),anchor="e",font=("Segoe UI",8))
        n=max(len(values),1); slot=pw/n; barw=max(4,slot*.58)
        for i,(lab,val) in enumerate(zip(labels,values)):
            x=left+slot*i+slot/2
            bh=ph*(val/maxv) if maxv else 0
            y=h-bottom-bh
            canvas.create_rectangle(x-barw/2,y,x+barw/2,h-bottom,fill="#4e79a7",outline="")
            if val: canvas.create_text(x,y-8,text=str(val),font=("Segoe UI",8,"bold"))
            canvas.create_text(x,h-bottom+16,text=lab,font=("Segoe UI",7),angle=45,anchor="ne")
        canvas.create_text(left+pw/2,10,text="Cantidad de movimientos",font=("Segoe UI",9,"bold"))

    def refresh_movement_charts(self):
        if not hasattr(self,"daily_chart"): return
        dates=self._movement_dates(); today=datetime.now().date()
        daily_days=[today-timedelta(days=i) for i in range(13,-1,-1)]
        daily_vals=[sum(1 for d in dates if d==day) for day in daily_days]
        daily_labels=[d.strftime("%d/%m") for d in daily_days]
        self._draw_bar_chart(self.daily_chart,daily_labels,daily_vals)

        this_monday=today-timedelta(days=today.weekday())
        starts=[this_monday-timedelta(weeks=i) for i in range(7,-1,-1)]
        weekly_vals=[sum(1 for d in dates if start<=d<start+timedelta(days=7)) for start in starts]
        weekly_labels=[start.strftime("%d/%m") for start in starts]
        self._draw_bar_chart(self.weekly_chart,weekly_labels,weekly_vals)

    def refresh_inventory(self):
        if not hasattr(self,"tree"):return
        q=self.search.get().lower(); [self.tree.delete(x) for x in self.tree.get_children()]
        wb=book()
        # La vista de Asignaciones cambia únicamente los encabezados/orden de presentación.
        # La estructura y los datos de las hojas Excel permanecen intactos.
        inv_by_id={}
        for ir in wb["Inventario"].iter_rows(min_row=2):
            key=str(ir[0].value or "").strip()
            if key: inv_by_id[key]=ir
        equip_by_imei={}
        if "Inventario Equipos" in wb.sheetnames:
            for er in wb["Inventario Equipos"].iter_rows(min_row=2):
                key=str(er[1].value or "").strip()
                if key: equip_by_imei[key]=er
        rows_added=[]
        for ar in wb["Asignaciones"].iter_rows(min_row=2):
            if str(ar[13].value or "").strip().casefold()!="activo": continue
            eqid=str(ar[1].value or "").strip()
            inv=inv_by_id.get(eqid)
            imei=str(ar[7].value or (inv[12].value if inv else "") or "").strip()
            equip=equip_by_imei.get(imei)
            visible=[
                ar[10].value if ar[10].value not in (None, "") else (inv[1].value if inv else ""), # Operador
                inv[2].value if inv else "", # Razón Social
                ar[2].value if ar[2].value not in (None, "") else (inv[3].value if inv else ""), # Nombre y Apellidos
                ar[4].value if ar[4].value not in (None, "") else (inv[4].value if inv else ""), # Correo
                ar[5].value if ar[5].value not in (None, "") else (inv[5].value if inv else ""), # Cargo
                ar[3].value if ar[3].value not in (None, "") else (inv[6].value if inv else ""), # DNI
                ar[9].value if ar[9].value not in (None, "") else (inv[7].value if inv else ""), # Número celular
                inv[8].value if inv else "", # Plan Celular
                inv[9].value if inv else "", # Tipo de Plan
                ar[6].value if ar[6].value not in (None, "") else (inv[10].value if inv else ""), # Modelo de Equipo
                inv[11].value if inv else "", # Tipo de Gama
                imei, # IMEI
                equip[5].value if equip else (inv[13].value if inv else "") # Situación
            ]
            if not q or q in " ".join(str(v or "") for v in visible).lower():
                self.tree.insert("", "end",values=visible)

    def refresh_equipment(self):
        if not hasattr(self,"et"):return
        [self.et.delete(x) for x in self.et.get_children()]
        for r in self.equipment_rows(): self.et.insert("", "end",values=r[1:])

    def refresh_history(self):
        [self.ht.delete(x) for x in self.ht.get_children()]
        movement_labels={
            "CARGA MASIVA":"Carga Masiva",
            "ASIGNACION":"Asignación",
            "MANTENIMIENTO":"Mantenimiento",
            "REHABILITACION":"Rehabilitación",
            "BAJA":"Baja",
            "ELIMINACION":"Eliminación",
            "RECUPERACION":"Recuperación",
        }
        for r in book()["Movimientos"].iter_rows(min_row=2,values_only=True):
            # ID Movimiento e ID Equipo no se muestran en la interfaz.
            raw_type=str(r[2] or "")
            display_type=movement_labels.get(raw_type,raw_type)
            detail=str(r[6] or "")
            # En Historial, ASIGNACION y ELIMINACION ya tienen una columna IMEI
            # independiente; por eso no se repite ni la palabra IMEI ni su valor
            # dentro de Detalle. Esto también normaliza registros históricos.
            if raw_type in ("ASIGNACION","ELIMINACION"):
                if raw_type == "ASIGNACION":
                    detail=detail.replace("IMEI ","")
                    import re
                    detail=re.sub(r"\bIMEI\s*[:=-]?\s*\d{14,20}\b", "", detail, flags=re.IGNORECASE)
                    # Para registros V20.3 del formato "IMEI 123... asignado a ..."
                    detail=re.sub(r"^\s*\d{14,20}\s+", "", detail)
                    detail=detail.replace("  "," ").strip()
                else:
                    import re
                    detail=re.sub(r"\s*\|?\s*IMEI\s*[:=-]?\s*\d{14,20}\b", "", detail, flags=re.IGNORECASE)
                    detail=detail.replace("  "," ").strip(" |")
            self.ht.insert("", "end",values=(display_type,r[3],r[4],r[5],detail))

    def refresh_actas_entrega(self):
        if not hasattr(self,"at"): return
        for x in self.at.get_children(): self.at.delete(x)
        if not EXCEL.exists(): return
        wb=book()
        if "Actas de Entrega" not in wb.sheetnames: return
        ws=wb["Actas de Entrega"]
        for r in ws.iter_rows(min_row=2,values_only=True):
            vals=list(r[:len(ACTAS_ENTREGA_HEADERS)])
            if not any(v not in (None,"") for v in vals): continue
            self.at.insert("","end",values=vals)

    def open_selected_acta(self,event=None):
        if not hasattr(self,"at"): return
        item=self.at.identify_row(event.y) if event is not None else (self.at.selection()[0] if self.at.selection() else "")
        if not item: return
        vals=list(self.at.item(item,"values"))
        if len(vals)<6: return
        filename=str(vals[5] or "").strip()
        if not filename:
            messagebox.showwarning("Acta de Entrega","El registro no tiene nombre de documento generado.")
            return
        pdf=ACTAS/filename
        if not pdf.exists():
            messagebox.showerror("Acta de Entrega",f"No se encuentra el archivo PDF:\n{pdf}")
            return
        try:
            if platform.system()=="Windows":
                os.startfile(str(pdf))
            elif platform.system()=="Darwin":
                subprocess.Popen(["open",str(pdf)])
            else:
                subprocess.Popen(["xdg-open",str(pdf)])
        except Exception as e:
            messagebox.showerror("Acta de Entrega",f"No fue posible abrir el PDF.\n\nDetalle: {e}")

    def refresh_users(self):
        if not hasattr(self,"ut"):return
        [self.ut.delete(x) for x in self.ut.get_children()]
        for r in book()["Usuarios"].iter_rows(min_row=2,values_only=True): self.ut.insert("", "end",values=(r[0],r[1]))

    def selected(self):
        s=self.tree.selection()
        if not s:messagebox.showwarning("Selección","Seleccione un registro.");return None
        return list(self.tree.item(s[0],"values"))

    def equipment_selected(self):
        s=self.et.selection()
        if not s:messagebox.showwarning("Selección","Seleccione un equipo.");return None
        return list(self.et.item(s[0],"values"))

    def find_inventory_by_imei(self,imei):
        wb=book()
        for r in wb["Inventario"].iter_rows(min_row=2):
            if str(r[12].value or "").strip()==str(imei).strip(): return r
        return None

    def equipment_data(self,imei):
        for r in self.equipment_rows():
            if str(r[1] or "").strip()==str(imei).strip(): return r
        return None

    def available_imeis(self):
        return [str(r[1]) for r in self.equipment_rows() if r[1] and str(r[5])=="Disponible"]

    def equipment_form(self):
        w=tk.Toplevel(self);w.title("Registrar equipo / IMEI");w.geometry("500x400")
        labels=["IMEI","Marca y Modelo","Gama Celular","Estado de equipo","Observaciones"];e={}
        for i,k in enumerate(labels):
            ttk.Label(w,text=k).grid(row=i,column=0,sticky="w",padx=10,pady=7)
            vals=GAMAS if k=="Gama Celular" else EQUIP_STATES if k=="Estado de equipo" else None
            e[k]=ttk.Combobox(w,values=vals,state="readonly",width=35) if vals else ttk.Entry(w,width=38)
            e[k].grid(row=i,column=1,padx=10,pady=7)
            if vals:e[k].set(vals[0])
        def save():
            imei=e["IMEI"].get().strip()
            if not imei.isdigit() or len(imei) not in (14,15):messagebox.showerror("IMEI","El IMEI debe tener 14 o 15 dígitos.");return
            wb=book();ws=wb["Inventario Equipos"]
            if any(str(r[1].value)==imei for r in ws.iter_rows(min_row=2)):messagebox.showerror("IMEI","El IMEI ya existe.");return
            iid=max([int(r[0].value or 0) for r in ws.iter_rows(min_row=2)],default=0)+1
            ws.append([iid,imei,e["Marca y Modelo"].get().strip(),e["Gama Celular"].get(),e["Estado de equipo"].get(),"Disponible",now(),e["Observaciones"].get().strip()])
            wb.save(EXCEL);log(self.user,iid,"ALTA EQUIPO","IMEI "+imei,imei);w.destroy();self.refresh_all()
        ttk.Button(w,text="Guardar",command=save).grid(row=6,column=1,pady=15)

    def assign(self):
        w=tk.Toplevel(self);w.title("Nueva asignación");w.geometry("520x500");es={}
        ttk.Label(w,text="IMEI").grid(row=0,column=0,padx=8,pady=7)
        ime_values=self.available_imeis()
        ime=ttk.Combobox(w,values=ime_values,state="normal",width=36)
        ime.grid(row=0,column=1)
        ime_hint=tk.StringVar(value="Escriba el IMEI completo. Se muestran coincidencias disponibles.")
        ttk.Label(w,textvariable=ime_hint,foreground="#555555").grid(row=0,column=2,padx=6,sticky="w")
        # Autocompletado no bloqueante: permite escribir los 14/15 dígitos completos.
        # El primer resultado siempre es la coincidencia disponible más cercana por prefijo.
        def update_imei_suggestions(event=None):
            typed="".join(ch for ch in ime.get().strip() if ch.isdigit())
            if typed != ime.get().strip():
                ime.delete(0,"end"); ime.insert(0,typed)
            if not typed:
                ime.configure(values=ime_values)
                ime_hint.set("Escriba el IMEI completo. Se muestran coincidencias disponibles.")
                model.set(""); gama.set(""); estate.set("")
                return
            matches=[x for x in ime_values if x.startswith(typed)]
            # Si no hay prefijo exacto, muestra los IMEI disponibles más cercanos por
            # distancia de edición y prefijo, sin impedir seguir escribiendo.
            if matches:
                matches=sorted(matches, key=lambda x:(len(x)-len(typed), x))
            else:
                def score(x):
                    common=0
                    for a,b in zip(typed,x):
                        if a!=b: break
                        common+=1
                    return (-common, abs(len(x)-len(typed)), x)
                matches=sorted(ime_values,key=score)
            ime.configure(values=matches)
            exact=self.equipment_data(typed) if typed in ime_values else None
            if exact:
                model.set(exact[2] or "");gama.set(exact[3] or "");estate.set(exact[4] or "")
                ime_hint.set(f"IMEI disponible: {typed} — {exact[2] or 'Sin marca/modelo'}")
            elif matches:
                closest=matches[0]
                cd=self.equipment_data(closest)
                ime_hint.set(f"Sugerencia: {closest} — {(cd[2] if cd else '') or 'Sin marca/modelo'}")
                model.set("");gama.set("");estate.set("")
            else:
                ime_hint.set("No hay IMEI disponibles.")
                model.set("");gama.set("");estate.set("")

        ttk.Label(w,text="Marca y Modelo").grid(row=1,column=0,padx=8,pady=7);model=tk.StringVar();ttk.Entry(w,textvariable=model,state="readonly",width=39).grid(row=1,column=1)
        ttk.Label(w,text="Gama").grid(row=2,column=0,padx=8,pady=7);gama=tk.StringVar();ttk.Entry(w,textvariable=gama,state="readonly",width=39).grid(row=2,column=1)
        ttk.Label(w,text="Estado de equipo").grid(row=3,column=0,padx=8,pady=7);estate=tk.StringVar();ttk.Entry(w,textvariable=estate,state="readonly",width=39).grid(row=3,column=1)
        fields=["Operador","Razon Social","Nombre y Apellido","Correo","Cargo","DNI","Número celular","Plan Celular","Tipo de Plan"]
        vals={};cats={"Operador":OPERADORES,"Razon Social":RAZONES,"Plan Celular":PLANES,"Tipo de Plan":TIPOS_PLAN}
        for i,k in enumerate(fields,4):
            ttk.Label(w,text=k).grid(row=i,column=0,padx=8,pady=5,sticky="w")
            vals[k]=ttk.Combobox(w,values=cats[k],state="readonly",width=36) if k in cats else ttk.Entry(w,width=39)
            vals[k].grid(row=i,column=1)
            if k in cats: vals[k].set(cats[k][0])
        def pick(*a):
            d=self.equipment_data(ime.get())
            if d:model.set(d[2] or "");gama.set(d[3] or "");estate.set(d[4] or "")
        ime.bind("<<ComboboxSelected>>",pick)
        ime.bind("<KeyRelease>",update_imei_suggestions)
        def save():
            if not ime.get() or not vals["Nombre y Apellido"].get().strip() or not vals["DNI"].get().strip():messagebox.showerror("Datos","Seleccione IMEI e indique Nombre y DNI.");return
            d=self.equipment_data(ime.get())
            if not d or str(d[5] or "").strip().casefold()!="disponible":messagebox.showerror("IMEI","El IMEI no está disponible.");return
            wb=book();wi=wb["Inventario Equipos"]
            for rr in wi.iter_rows(min_row=2):
                if str(rr[1].value)==ime.get():rr[5].value="Asignado"
            ws=wb["Inventario"];eq=max([int(r[0].value or 0) for r in ws.iter_rows(min_row=2)],default=0)+1
            row=[eq,vals["Operador"].get(),vals["Razon Social"].get(),vals["Nombre y Apellido"].get().strip(),vals["Correo"].get().strip(),vals["Cargo"].get().strip(),vals["DNI"].get().strip(),vals["Número celular"].get().strip(),vals["Plan Celular"].get(),vals["Tipo de Plan"].get(),d[2],d[3],ime.get(),"Asignado",d[4],now(),"",""]
            ws.append(row)
            wa=wb["Asignaciones"];aid=max([int(r[0].value or 0) for r in wa.iter_rows(min_row=2)],default=0)+1
            wa.append([aid,eq,row[3],row[6],row[4],row[5],row[10],row[12],row[14],row[7],row[1],now(),"","Activo","Asignación desde Inventario Equipos / IMEI"])
            wb.save(EXCEL);log(self.user,eq,"ASIGNACION","Asignado a "+row[3],ime.get());w.destroy();self.refresh_all()
        ttk.Button(w,text="Guardar asignación",command=save).grid(row=13,column=1,pady=15)

    def release(self):
        # "Recuperar" solicita una observación obligatoria antes de liberar el equipo.
        obs=simpledialog.askstring("Recuperar equipo","Indique la observación / motivo de recuperación del equipo:",parent=self)
        if obs is None or not obs.strip():
            messagebox.showwarning("Observación","Debe ingresar una observación para recuperar el equipo.");return
        self.change_state("Disponible",release=True,observation=obs.strip())

    def maintenance(self):
        # Flujo de Mantenimiento: primero solicita una observación obligatoria;
        # solo después cambia el equipo a Mantenimiento y elimina su asignación activa.
        obs=simpledialog.askstring(
            "Enviar a mantenimiento",
            "Indique la observación / motivo de mantenimiento del equipo antes de liberarlo:",
            parent=self
        )
        if obs is None or not obs.strip():
            messagebox.showwarning(
                "Observación",
                "Debe ingresar una observación para enviar el equipo a mantenimiento. El equipo no será liberado ni se eliminará la asignación.",
                parent=self
            )
            return
        self.change_state("Mantenimiento",observation=obs.strip())

    def baja(self):
        r=self.selected()
        if not r:return
        w=tk.Toplevel(self);w.title("Dar de baja equipo");w.resizable(False,False);w.transient(self);w.grab_set()
        ttk.Label(w,text="Observación de baja").grid(row=0,column=0,padx=12,pady=12,sticky="w")
        reason=tk.StringVar()
        cb=ttk.Combobox(w,textvariable=reason,values=["Robo","Pérdida","Destrucción"],state="readonly",width=24)
        # Baja usa únicamente esta observación predefinida; no se solicita texto libre.
        cb.grid(row=0,column=1,padx=12,pady=12);cb.current(0)
        def save_baja():
            full=reason.get().strip()
            if not full:
                messagebox.showwarning("Motivo","Seleccione el motivo de baja.",parent=w);return
            w.destroy();self.change_state("Baja",observation=full)
        ttk.Button(w,text="Confirmar baja",command=save_baja).grid(row=1,column=1,pady=14,sticky="e")
        cb.focus_set()

    def change_state(self,state,release=False,observation=""):
        r=self.selected()
        if not r:return
        imei=str(r[11]) # vista Asignaciones: IMEI es el índice 11
        inv=self.find_inventory_by_imei(imei)
        if inv is None: messagebox.showerror("Equipo","No se encontró el IMEI en asignaciones.");return
        current_state=str(inv[13].value or "").strip()
        # Un equipo en Baja es definitivo: no puede volver a Disponible, Asignado
        # ni Mantenimiento desde ningún flujo de cambio de estado.
        if current_state.casefold()=="baja" and state.casefold()!="baja":
            messagebox.showwarning("Equipo en Baja","El equipo está dado de Baja y no puede ser rehabilitado ni cambiar de estado.",parent=self)
            return
        wb=book()
        obs=observation.strip()
        eqid=str(inv[0].value)
        # Cambiar el estado del equipo y conservar la observación en el inventario.
        # Regla V20.1: al recuperar un equipo o enviarlo a mantenimiento,
        # su "Estado de equipo" pasa de "Equipo nuevo" a "Equipo seminuevo".
        # Esto se actualiza tanto en Inventario como en Inventario Equipos para
        # mantener ambas hojas consistentes y para que el Acta posterior muestre
        # correctamente el estado físico del equipo.
        mark_as_seminew = release or state.casefold() == "mantenimiento"
        for rr in wb["Inventario"].iter_rows(min_row=2):
            if str(rr[12].value)==imei:
                rr[13].value=state
                if mark_as_seminew:
                    rr[14].value="Equipo seminuevo"
                rr[16].value=now()
                if obs: rr[17].value=obs
        for rr in wb["Inventario Equipos"].iter_rows(min_row=2):
            if str(rr[1].value)==imei:
                rr[5].value=state
                if mark_as_seminew:
                    rr[4].value="Equipo seminuevo"
                if obs: rr[7].value=obs
        # Recuperar, Mantenimiento y Baja eliminan físicamente el registro activo
        # de la pestaña Asignaciones. En Mantenimiento esto ocurre únicamente
        # después de haber recibido la observación obligatoria.
        wa=wb["Asignaciones"]
        # La eliminación de la asignación se hace por IMEI y/o ID Equipo para
        # evitar que una variación del ID impida retirar el registro visible.
        # Se eliminan todas las asignaciones activas que correspondan al equipo.
        deleted_assignments=0
        for i in range(wa.max_row,1,-1):
            rr=wa[i]
            row_eqid=str(rr[1].value or "").strip()
            row_imei=str(rr[7].value or "").strip()
            row_status=str(rr[13].value or "").strip().casefold()
            if row_status == "activo" and (row_eqid == eqid or row_imei == imei):
                wa.delete_rows(i,1)
                deleted_assignments += 1
        if deleted_assignments == 0:
            # Respaldo para bases antiguas donde el estado de Asignaciones
            # puede no estar normalizado, pero el IMEI sí identifica la fila.
            for i in range(wa.max_row,1,-1):
                rr=wa[i]
                row_eqid=str(rr[1].value or "").strip()
                row_imei=str(rr[7].value or "").strip()
                if row_eqid == eqid or row_imei == imei:
                    wa.delete_rows(i,1)
                    deleted_assignments += 1
        wb.save(EXCEL)
        action="RECUPERACION" if release else ("MANTENIMIENTO" if state=="Mantenimiento" else "BAJA" if state=="Baja" else "CAMBIO ESTADO")
        detail=state+(" | "+obs if obs else "")
        log(self.user,int(inv[0].value),action,detail,imei);self.refresh_all()

    def delete(self):
        r=self.selected()
        if not r:return
        imei=str(r[11]).strip()
        inv=self.find_inventory_by_imei(imei)
        if not messagebox.askyesno("Eliminar","¿Eliminar el registro seleccionado de Asignaciones?",parent=self):
            return
        wb=book();wa=wb["Asignaciones"]
        deleted=0;eqid=str(inv[0].value).strip() if inv is not None else ""
        # El botón Eliminar debe borrar la fila de Asignaciones, no la fila de
        # Inventario. El inventario sigue siendo necesario para controlar el equipo.
        for i in range(wa.max_row,1,-1):
            rr=wa[i]
            row_eqid=str(rr[1].value or "").strip()
            row_imei=str(rr[7].value or "").strip()
            if (row_imei == imei or (eqid and row_eqid == eqid)):
                wa.delete_rows(i,1)
                deleted += 1
        if deleted == 0:
            messagebox.showwarning("Eliminar","No se encontró el registro seleccionado en Asignaciones.",parent=self)
            return
        # Al eliminar manualmente una asignación, el equipo vuelve a estar
        # disponible para una nueva asignación.
        for rr in wb["Inventario"].iter_rows(min_row=2):
            if str(rr[12].value or "").strip() == imei:
                rr[13].value="Disponible"
                rr[16].value=now()
        if "Inventario Equipos" in wb.sheetnames:
            for rr in wb["Inventario Equipos"].iter_rows(min_row=2):
                if str(rr[1].value or "").strip() == imei:
                    rr[5].value="Disponible"
                    rr[7].value="Eliminación manual de asignación"
        wb.save(EXCEL)
        log(self.user,int(inv[0].value) if inv is not None else 0,"ELIMINACION","Asignación eliminada",imei)
        self.refresh_all()

    def bulk(self):
        p=filedialog.askopenfilename(filetypes=[("Excel","*.xlsx")])
        if not p:return
        try: src=load_workbook(p,data_only=True);sw=src.active;hdr=[str(c.value or "").strip() for c in sw[1]]
        except Exception as e: messagebox.showerror("Carga masiva",f"No se pudo leer el Excel.\n\n{e}");return
        required=["IMEI"]
        if "Marca y Modelo" not in hdr and "Modelo Asignado" not in hdr and not ("Marca" in hdr and "Modelo" in hdr):
            messagebox.showerror("Carga masiva","El Excel debe contener las columnas: IMEI y Marca y Modelo.");return
        wb=book();ws=wb["Inventario Equipos"];existing={str(r[1].value) for r in ws.iter_rows(min_row=2) if r[1].value};count=0;skipped=0
        nextid=max([int(r[0].value or 0) for r in ws.iter_rows(min_row=2)],default=0)+1
        for vals in sw.iter_rows(min_row=2,values_only=True):
            d=dict(zip(hdr,vals)); raw=d.get("IMEI",""); imei=str(raw).strip()
            if imei.endswith(".0"): imei=imei[:-2]
            if not imei.isdigit() or len(imei) not in (14,15) or imei in existing: skipped+=1;continue
            combined=d.get("Marca y Modelo","") or d.get("Modelo Asignado","") or " ".join(x for x in (str(d.get("Marca","") or "").strip(),str(d.get("Modelo","") or "").strip()) if x)
            ws.append([nextid,imei,combined,d.get("Gama Celular","NA") or "NA",d.get("Estado de equipo","Equipo nuevo") or "Equipo nuevo","Disponible",now(),"Carga masiva"])
            existing.add(imei);count+=1;nextid+=1
        wb.save(EXCEL)
        log(self.user,"","CARGA MASIVA",f"Equipos importados: {count}; omitidos: {skipped}")
        messagebox.showinfo("Carga masiva",f"Se importaron {count} equipos/IMEI.\nRegistros omitidos: {skipped}.\n\nLa Situación de los equipos importados queda como Disponible.")
        self.refresh_all()

    def import_help(self):
        messagebox.showinfo("Importación Excel",
            "FORMATO DEL EXCEL IMPORTADOR\n\n"
            "1. El archivo debe ser .xlsx.\n"
            "2. La primera fila debe contener los encabezados.\n"
            "3. Campos obligatorios:\n   • IMEI\n   • Marca y Modelo\n"
            "4. Campos opcionales:\n   • Gama Celular\n   • Estado de equipo (Equipo nuevo / Equipo seminuevo)\n"
            "5. Cada IMEI debe tener 14 o 15 dígitos y no repetirse.\n"
            "6. Los equipos importados ingresan con Situación = Disponible.\n\n"
            "Ejemplo de encabezados:\nIMEI | Marca y Modelo | Gama Celular | Estado de equipo\n\n"
            "No es necesario incluir ID, Situación, Fecha alta ni Observaciones; el sistema los genera automáticamente.")

    def rehabilitate(self):
        r=self.equipment_selected()
        if not r:return
        imei=str(r[0]); situation=str(r[4])
        if situation == "Baja":
            messagebox.showwarning("Rehabilitar","Un equipo dado de Baja no puede ser rehabilitado.");return
        if situation != "Mantenimiento":
            messagebox.showwarning("Rehabilitar","Solo se pueden rehabilitar equipos en Mantenimiento.");return
        obs=simpledialog.askstring("Rehabilitar equipo","Indique la observación / motivo de rehabilitación:",parent=self)
        if obs is None or not obs.strip():
            messagebox.showwarning("Observación","Debe ingresar una observación para rehabilitar el equipo.");return
        wb=book();ws=wb["Inventario Equipos"]
        for rr in ws.iter_rows(min_row=2):
            if str(rr[1].value)==imei:
                rr[5].value="Disponible"; rr[7].value=obs.strip();break
        wb.save(EXCEL);log(self.user,imei,"REHABILITACION","Equipo rehabilitado: "+obs.strip(),imei);self.refresh_all()

    def export_xlsx(self):
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",filetypes=[("Excel","*.xlsx")],initialfile="Gestion_Celulares_Exportacion.xlsx")
        if not p:return
        src=book(); out=Workbook()
        default=out.active; out.remove(default)

        # El XLSX exportado es un reporte: no expone identificadores internos.
        hidden_by_sheet={
            "Inventario":{"id"},
            "Inventario Equipos":{"id equipo"},
            "Asignaciones":{"id asignación","id equipo"},
            "Movimientos":{"id movimiento","id equipo"},
            "Actas":{"id equipo"},
        }
        allowed=["Inventario","Inventario Equipos","Asignaciones"]
        if self.role=="Administrador": allowed += ["Movimientos","Usuarios","Actas"]

        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        header_fill=PatternFill("solid", fgColor="1F4E78")
        header_font=Font(bold=True, color="FFFFFF")
        thin=Side(style="thin", color="D9E2F3")
        header_border=Border(bottom=thin)

        for name in allowed:
            sw=src[name]; dw=out.create_sheet(name)
            headers=[str(c.value or "").strip() for c in sw[1]]
            omit=hidden_by_sheet.get(name,set())
            keep=[i for i,h in enumerate(headers) if h.casefold() not in omit]

            # Copia solo las columnas permitidas.
            for out_col, src_idx in enumerate(keep, start=1):
                src_cell=sw.cell(1,src_idx+1)
                dst=dw.cell(1,out_col,src_cell.value)
                dst.font=header_font; dst.fill=header_fill; dst.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); dst.border=header_border
            for src_row in range(2,sw.max_row+1):
                for out_col, src_idx in enumerate(keep, start=1):
                    src_cell=sw.cell(src_row,src_idx+1)
                    dst=dw.cell(src_row,out_col,src_cell.value)
                    dst.number_format=src_cell.number_format

            dw.freeze_panes="A2"
            dw.auto_filter.ref=dw.dimensions if dw.max_row >= 1 and dw.max_column >= 1 else None
            dw.row_dimensions[1].height=30
            for col in dw.columns:
                letter=col[0].column_letter
                maxlen=max((len(str(c.value or "")) for c in col),default=0)
                dw.column_dimensions[letter].width=min(max(maxlen+2,12),35)

            # Fechas del reporte.
            for cell in dw[1]:
                h=str(cell.value or "").casefold()
                if "fecha" in h:
                    for row in range(2,dw.max_row+1):
                        dw.cell(row,cell.column).number_format="dd/mm/yyyy"

        out.save(p)
        sheets=", ".join(allowed)
        messagebox.showinfo("Exportación",f"Archivo XLSX generado correctamente.\n\nHojas incluidas: {sheets}\n\nLos identificadores internos fueron excluidos del archivo exportado.")

    def acta(self):
        r=self.selected()
        if not r:return

        def clean_value(value):
            if value is None:
                return ""
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value).strip()

        # El IMEI y el celular se toman primero de la fila visible seleccionada.
        imei=clean_value(r[11]) if len(r)>11 else ""
        celular_visible=clean_value(r[6]) if len(r)>6 else ""
        if not imei:
            messagebox.showwarning("Acta","El registro seleccionado no tiene IMEI.");return

        wb=book()
        wa=wb["Asignaciones"]

        def header_map(ws):
            return {str(c.value or "").strip().casefold(): i for i,c in enumerate(ws[1])}

        ca=header_map(wa)
        modelo_col = "marca y modelo" if "marca y modelo" in ca else ("modelo asignado" if "modelo asignado" in ca else None)
        required_a=["imei","número celular","nombre y apellido","cargo","dni","operador","estado","id equipo"]
        missing_a=[h for h in required_a if h not in ca]
        if modelo_col is None:
            missing_a.append("marca y modelo")
        if missing_a:
            messagebox.showerror("Acta","Faltan columnas requeridas en 'Asignaciones':\n"+", ".join(missing_a));return

        asign=None
        for ar in wa.iter_rows(min_row=2):
            if clean_value(ar[ca["imei"]].value)==imei and clean_value(ar[ca["estado"]].value).casefold()=="activo":
                asign=ar; break
        if asign is None:
            messagebox.showerror("Acta","No se encontró una asignación activa para el IMEI seleccionado en la hoja 'Asignaciones'.")
            return

        # Primero buscamos en Inventario Equipos. Si la base todavía no tiene
        # esa fila, usamos el registro equivalente de Inventario (por IMEI o ID).
        equipo=None
        estado_equipo=""
        if "Inventario Equipos" in wb.sheetnames:
            we=wb["Inventario Equipos"]
            ce=header_map(we)
            if "imei" in ce and "estado de equipo" in ce:
                for er in we.iter_rows(min_row=2):
                    if clean_value(er[ce["imei"]].value)==imei:
                        equipo=er; estado_equipo=clean_value(er[ce["estado de equipo"]].value); break
                if equipo is None and "id equipo" in ce:
                    eqid=clean_value(asign[ca["id equipo"]].value)
                    if eqid:
                        for er in we.iter_rows(min_row=2):
                            if clean_value(er[ce["id equipo"]].value)==eqid:
                                equipo=er; estado_equipo=clean_value(er[ce["estado de equipo"]].value); break

        inv_match=None
        if "Inventario" in wb.sheetnames:
            wi=wb["Inventario"]
            ci=header_map(wi)
            # El Excel proporcionado usa ID como identificador del equipo.
            eqid=clean_value(asign[ca["id equipo"]].value)
            for ir in wi.iter_rows(min_row=2):
                if ("imei" in ci and clean_value(ir[ci["imei"]].value)==imei) or ("id" in ci and eqid and clean_value(ir[ci["id"]].value)==eqid):
                    # Preferimos el registro Asignado cuando hay duplicados.
                    if inv_match is None or clean_value(ir[ci.get("estado",0)].value).casefold()=="asignado":
                        inv_match=ir
            if not estado_equipo and inv_match is not None and "estado de equipo" in ci:
                estado_equipo=clean_value(inv_match[ci["estado de equipo"]].value)

        if not estado_equipo:
            messagebox.showerror("Acta","No se encontró el 'Estado de equipo' para el IMEI seleccionado.")
            return

        nombre=clean_value(asign[ca["nombre y apellido"]].value)
        cargo=clean_value(asign[ca["cargo"]].value)
        dni=clean_value(asign[ca["dni"]].value)
        modelo=clean_value(asign[ca[modelo_col]].value)
        imei_db=clean_value(asign[ca["imei"]].value)
        celular_db=clean_value(asign[ca["número celular"]].value)
        if inv_match is not None:
            ci=header_map(wb["Inventario"])
            if "imei" in ci: imei_db=imei_db or clean_value(inv_match[ci["imei"]].value)
            if "número celular" in ci: celular_db=celular_db or clean_value(inv_match[ci["número celular"]].value)
            if not modelo and "modelo asignado" in ci: modelo=clean_value(inv_match[ci["modelo asignado"]].value)
        imei_acta=imei_db or imei
        celular=celular_db or celular_visible
        operador=clean_value(asign[ca["operador"]].value)
        fecha_entrega=datetime.now().strftime("%d/%m/%Y")

        if not nombre or not dni or not imei_acta:
            messagebox.showwarning("Acta","La asignación debe contener Nombre y Apellidos, DNI e IMEI.");return
        if not celular:
            messagebox.showwarning("Acta","La asignación seleccionada no tiene Número celular. Complete el campo en Asignaciones y vuelva a generar el acta.");return
        # La plantilla DOCX es externa; el script PS1 sí permanece integrado y
        # se materializa temporalmente durante la generación.

        n=max([int(float(x[0].value or 0)) for x in wb["Actas"].iter_rows(min_row=2) if str(x[0].value or "").strip().replace(".","",1).isdigit()],default=0)+1
        nombre_seguro="".join(c for c in nombre if c not in '<>:/\\|?*').strip() or "Sin Nombre"
        pdf=ACTAS/f"ACTA DE ENTREGA - {nombre_seguro}.pdf"
        repl={
            "{{Nombre y Apellidos}}":nombre,
            "{{Cargo}}":cargo,
            "{{DNI}}":dni,
            "{{Modelo Asignado}}":modelo,
            "{{Número de IMEI}}":imei_acta,
            "{{Estado de equipo}}":estado_equipo,
            "{{Número celular}}":celular,
            "{{Operador}}":operador,
            "{{Fecha de entrega}}":fecha_entrega
        }

        # La plantilla DOCX permanece externa. Se crea una copia temporal y se
        # reemplazan directamente los marcadores en el XML interno del DOCX.
        # Esto evita que Word Find falle con los marcadores de IMEI y celular
        # cuando están dentro de celdas de tablas. La plantilla original nunca
        # se modifica ni se guarda.
        temp_ps=ACTAS/f".generar_acta_word_{n:05d}.ps1"
        payload=ACTAS/f".acta_{n:05d}.json"
        temp_docx=ACTAS/f".acta_template_{n:05d}.docx"
        try:
            materialize_acta_resources(temp_ps)
            if not TEMPLATE.exists():
                raise FileNotFoundError(f"No se encontró la plantilla: {TEMPLATE.name}")
            payload.write_text(json.dumps(repl,ensure_ascii=False),encoding="utf-8")

            with zipfile.ZipFile(TEMPLATE,"r") as zin, zipfile.ZipFile(temp_docx,"w",zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data=zin.read(item.filename)
                    if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                        try:
                            text=data.decode("utf-8")
                            for marker,value in repl.items():
                                text=text.replace(marker,str(value or ""))
                            data=text.encode("utf-8")
                        except UnicodeDecodeError:
                            pass
                    zout.writestr(item,data)

            # Validación previa: los nueve marcadores deben haber desaparecido
            # de todos los XML del DOCX temporal antes de llamar a Word.
            with zipfile.ZipFile(temp_docx,"r") as zcheck:
                remaining=[]
                for item in zcheck.infolist():
                    if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                        text=zcheck.read(item.filename).decode("utf-8",errors="ignore")
                        for marker in repl:
                            if marker in text:
                                remaining.append(marker)
                if remaining:
                    raise RuntimeError("No se pudieron reemplazar todos los marcadores del acta: " + ", ".join(sorted(set(remaining))))

            if platform.system()!="Windows": raise RuntimeError("La generación del acta requiere Windows con Microsoft Word.")
            subprocess.run(["powershell.exe","-NoProfile","-ExecutionPolicy","Bypass","-File",str(temp_ps),"-Template",str(temp_docx),"-OutputPdf",str(pdf),"-DataJson",str(payload)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True)
            if not pdf.exists(): raise RuntimeError("Microsoft Word no generó el PDF esperado.")
        except Exception as e:
            detail=(e.stderr if hasattr(e,"stderr") and e.stderr else str(e))
            messagebox.showerror("Acta PDF",f"No fue posible generar el PDF con Microsoft Word.\n\nVerifique que Microsoft Word esté instalado y activado.\n\nDetalle: {detail}");return
        finally:
            for f in (payload,temp_ps,temp_docx):
                try: f.unlink()
                except Exception: pass

        wb=book()
        fecha_generacion=datetime.now().strftime("%d/%m/%Y")
        wb["Actas"].append([n,asign[ca["id equipo"]].value,"Entrega",nombre,dni,fecha_generacion,pdf.name])
        if "Actas de Entrega" not in wb.sheetnames:
            ws_ae=wb.create_sheet("Actas de Entrega")
            ws_ae.append(ACTAS_ENTREGA_HEADERS)
        wb["Actas de Entrega"].append([fecha_generacion,self.user,nombre,dni,fecha_entrega,pdf.name])
        wb.save(EXCEL)
        self.refresh_actas_entrega()
        messagebox.showinfo("Acta PDF",f"Acta generada correctamente:\n{pdf}")

    def create_user(self):
        u=self.eu.get().strip();p=self.ep.get();role=self.er.get()
        if not u or not p or role not in ROLES:messagebox.showerror("Usuario","Complete usuario, clave y rol.");return
        wb=book();ws=wb["Usuarios"]
        if any(r[0].value==u for r in ws.iter_rows(min_row=2)):messagebox.showerror("Usuario","El usuario ya existe.");return
        ws.append([u,role,sha(p)]);wb.save(EXCEL);log(self.user,"","USUARIO","Creado "+u);self.eu.delete(0,"end");self.ep.delete(0,"end");self.refresh_users()
    def delete_user(self):
        s=self.ut.selection()
        if not s:return
        u=self.ut.item(s[0],"values")[0]
        if u==self.user or u=="admin":messagebox.showwarning("Usuario","No puede eliminarse ese usuario.");return
        wb=book();ws=wb["Usuarios"]
        for i in range(2,ws.max_row+1):
            if ws.cell(i,1).value==u:ws.delete_rows(i);break
        wb.save(EXCEL);log(self.user,"","USUARIO","Eliminado "+u);self.refresh_users()

if __name__=="__main__":
    try:
        ensure_equipment_schema()
        ensure_assignments_schema()
        ensure_actas_entrega_schema()
        normalize_application_dates()
        format_database_headers()
    except Exception as e:
        try:
            messagebox.showwarning("Base de datos", f"No se pudo completar la preparación de la base de datos.\n\n{e}")
        except Exception:
            pass
    Login().mainloop()
