from __future__ import annotations

import json
import os
import platform
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path

from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from werkzeug.utils import secure_filename


BASE = Path(__file__).resolve().parent
DATA_BASE = Path(os.environ.get("SIGE_DATA_DIR", BASE))
EXCEL = DATA_BASE / "Base_Datos_Celulares.xlsx"
ACTAS = DATA_BASE / "Actas Generadas"
TEMPLATE = DATA_BASE / "Acta de Entrega.docx"
SOURCE_EXCEL = BASE / "Base_Datos_Celulares.xlsx"
SOURCE_TEMPLATE = BASE / "Acta de Entrega.docx"

STATES = ["Disponible", "Asignado", "Mantenimiento", "Baja"]
EQUIP_STATES = ["Equipo nuevo", "Equipo seminuevo"]
OPERADORES = ["Claro", "Movistar", "Entel", "Bitel"]
RAZONES = ["Inchcape Motors Peru S.A.", "Inchcape Automotriz Peru S.A.", "Derco Peru S.A.", "Dercocenter S.A.C"]
PLANES = ["31.90", "44.90", "52.90", "55.90", "59.90", "62.90", "79.90", "98.90", "99.90", "105.90", "109.90", "119.90", "129.90", "199.90", "201.50", "249.00", "NA"]
TIPOS_PLAN = ["Plan 1", "Plan 2 y 3", "Plan 4", "BAM", "Chip", "NA"]
GAMAS = ["Gama 1", "Gama 2", "Gama 3", "Gama 4", "Gama 5", "NA"]
ROLES = ["Administrador", "Usuario Asignador", "Consulta"]
ACTAS_ENTREGA_HEADERS = ["Fecha", "Usuario", "Nombre", "DNI", "Fecha de Entrega", "Nombre del documento generado"]

app = Flask(
    __name__,
    template_folder=str(BASE / "templates_web"),
    static_folder=str(BASE / "static_web"),
)
app.config["SECRET_KEY"] = os.environ.get("SIGE_SECRET_KEY", "cambie-esta-clave-en-produccion")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024


def sha(s):
    import hashlib

    return hashlib.sha256(s.encode()).hexdigest()


def now():
    return datetime.now().strftime("%d/%m/%Y")


def book():
    return load_workbook(EXCEL)


def fmt_date(value):
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%d/%m/%Y")
        except Exception:
            pass
    try:
        return datetime.fromisoformat(text).strftime("%d/%m/%Y")
    except Exception:
        return text.split(" ")[0]


def initialize_runtime_files():
    """Prepara archivos editables para runtimes serverless como Vercel."""
    EXCEL.parent.mkdir(parents=True, exist_ok=True)
    ACTAS.mkdir(parents=True, exist_ok=True)
    if not EXCEL.exists() and SOURCE_EXCEL.exists():
        shutil.copy2(SOURCE_EXCEL, EXCEL)
    if not TEMPLATE.exists() and SOURCE_TEMPLATE.exists():
        shutil.copy2(SOURCE_TEMPLATE, TEMPLATE)


def ensure_sheet(wb, name, headers):
    if name not in wb.sheetnames:
        ws = wb.create_sheet(name)
        ws.append(headers)
        return ws
    ws = wb[name]
    if ws.max_row < 1:
        ws.append(headers)
    return ws


def ensure_assignments_schema():
    if not EXCEL.exists():
        return
    wb = load_workbook(EXCEL)
    headers = ["ID Asignacion", "ID Equipo", "Nombre y Apellido", "DNI", "Correo", "Cargo", "Marca y Modelo", "IMEI", "Estado de equipo", "Numero celular", "Operador", "Fecha Asignacion", "Fecha Liberacion", "Estado", "Detalle"]
    accented = ["ID Asignación", "ID Equipo", "Nombre y Apellido", "DNI", "Correo", "Cargo", "Marca y Modelo", "IMEI", "Estado de equipo", "Número celular", "Operador", "Fecha Asignación", "Fecha Liberación", "Estado", "Detalle"]
    ws = ensure_sheet(wb, "Asignaciones", accented)
    current = [clean(cell.value) for cell in ws[1]]
    if not any(current):
        ws.delete_rows(1, ws.max_row)
        ws.append(accented)
    elif len(current) < len(headers):
        for col in range(len(current) + 1, len(accented) + 1):
            ws.cell(1, col).value = accented[col - 1]
    wb.save(EXCEL)


def ensure_equipment_schema():
    if not EXCEL.exists():
        return
    wb = load_workbook(EXCEL)
    ensure_sheet(wb, "Inventario Equipos", ["ID Equipo", "IMEI", "Marca y Modelo", "Gama Celular", "Estado de equipo", "Situacion", "Fecha alta", "Observaciones"])
    wb.save(EXCEL)


def ensure_movimientos_schema():
    if not EXCEL.exists():
        return
    wb = load_workbook(EXCEL)
    ensure_sheet(wb, "Movimientos", ["ID Movimiento", "ID Equipo", "Tipo Movimiento", "Fecha", "Usuario", "IMEI", "Detalle"])
    wb.save(EXCEL)


def ensure_actas_entrega_schema():
    if not EXCEL.exists():
        return
    wb = load_workbook(EXCEL)
    ensure_sheet(wb, "Actas de Entrega", ACTAS_ENTREGA_HEADERS)
    ensure_sheet(wb, "Actas", ["ID Acta", "ID Equipo", "Tipo", "Nombre", "DNI", "Fecha", "Archivo PDF"])
    ensure_sheet(wb, "Actas PDF", ["Nombre del documento generado", "Contenido Base64"])
    wb.save(EXCEL)


def normalize_application_dates():
    if not EXCEL.exists():
        return
    wb = load_workbook(EXCEL)
    mappings = {
        "Inventario": [15, 16],
        "Inventario Equipos": [6],
        "Asignaciones": [11, 12],
        "Movimientos": [3],
        "Actas": [5],
        "Actas de Entrega": [0, 4],
    }
    changed = False
    for sheet, indexes in mappings.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2):
            for idx in indexes:
                if idx >= len(row):
                    continue
                cell = row[idx]
                if cell.value not in (None, ""):
                    value = fmt_date(cell.value)
                    if value and str(cell.value) != value:
                        cell.value = value
                        changed = True
    if changed:
        wb.save(EXCEL)


def format_database_headers():
    if not EXCEL.exists():
        return
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = load_workbook(EXCEL)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF")
    border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for ws in wb.worksheets:
        for cell in ws[1]:
            if cell.value is not None:
                cell.font = font
                cell.fill = fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
        ws.freeze_panes = "A2"
        if ws.max_row >= 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(EXCEL)


def log(user, eq, typ, detail, imei=None):
    wb = book()
    ws = ensure_sheet(wb, "Movimientos", ["ID Movimiento", "ID Equipo", "Tipo Movimiento", "Fecha", "Usuario", "IMEI", "Detalle"])
    if not imei:
        imei = ""
    n = max([int_or_zero(row[0].value) for row in ws.iter_rows(min_row=2)], default=0) + 1
    ws.append([n, eq, typ, now(), user, imei, detail])
    wb.save(EXCEL)


def save_pdf_backup(filename: str, pdf_path: Path):
    if not pdf_path.exists():
        return
    import base64

    wb = book()
    ws = ensure_sheet(wb, "Actas PDF", ["Nombre del documento generado", "Contenido Base64"])
    encoded = base64.b64encode(pdf_path.read_bytes()).decode("ascii")
    for row in ws.iter_rows(min_row=2):
        if clean(row[0].value) == filename:
            row[1].value = encoded
            wb.save(EXCEL)
            return
    ws.append([filename, encoded])
    wb.save(EXCEL)


def restore_pdf_from_backup(filename: str) -> Path | None:
    import base64

    if not EXCEL.exists():
        return None
    wb = book()
    if "Actas PDF" not in wb.sheetnames:
        return None
    for row in wb["Actas PDF"].iter_rows(min_row=2, values_only=True):
        if clean(row[0]) == filename and clean(row[1]):
            path = ACTAS / filename
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(base64.b64decode(clean(row[1])))
            return path
    return None


def prepare_database():
    initialize_runtime_files()
    ensure_assignments_schema()
    ensure_equipment_schema()
    ensure_actas_entrega_schema()
    ensure_movimientos_schema()
    normalize_application_dates()
    format_database_headers()


def current_user():
    if "user" not in session:
        return None
    return {"name": session["user"], "role": session.get("role", "Consulta")}


def level(role: str | None = None) -> int:
    role = role or session.get("role", "Consulta")
    return {"Consulta": 1, "Usuario Asignador": 2, "Administrador": 3}.get(role, 1)


def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not current_user():
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapper


def role_required(min_level: int):
    def decorator(view):
        @wraps(view)
        def wrapper(*args, **kwargs):
            if level() < min_level:
                flash("No tiene permisos para realizar esta accion.", "error")
                return redirect(url_for("dashboard"))
            return view(*args, **kwargs)

        return wrapper

    return decorator


def values_from_sheet(ws):
    return [list(row) for row in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in row)]


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def int_or_zero(value) -> int:
    try:
        return int(float(value or 0))
    except Exception:
        return 0


def header_map(ws):
    return {clean(cell.value).casefold(): i for i, cell in enumerate(ws[1])}


def sheet_rows(name: str):
    wb = book()
    if name not in wb.sheetnames:
        return []
    return values_from_sheet(wb[name])


def dashboard_counts():
    rows = sheet_rows("Inventario Equipos")
    counts = {"Total equipos": len(rows)}
    for state in STATES:
        counts[state] = sum(1 for row in rows if clean(row[5] if len(row) > 5 else "") == state)
    return counts


def movement_dates():
    dates = []
    wb = book()
    if "Movimientos" not in wb.sheetnames:
        return dates
    for row in wb["Movimientos"].iter_rows(min_row=2, values_only=True):
        value = row[3] if len(row) > 3 else None
        if isinstance(value, datetime):
            dates.append(value.date())
            continue
        if value in (None, ""):
            continue
        text = str(value).strip().split()[0]
        parsed = None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                parsed = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                pass
        if parsed:
            dates.append(parsed)
    return dates


def movement_chart_data():
    dates = movement_dates()
    today = datetime.now().date()

    daily_days = [today - timedelta(days=i) for i in range(13, -1, -1)]
    daily_values = [sum(1 for date in dates if date == day) for day in daily_days]

    this_monday = today - timedelta(days=today.weekday())
    week_starts = [this_monday - timedelta(weeks=i) for i in range(7, -1, -1)]
    weekly_values = [
        sum(1 for date in dates if start <= date < start + timedelta(days=7))
        for start in week_starts
    ]

    daily_max = max(daily_values + [1])
    weekly_max = max(weekly_values + [1])
    return {
        "daily": [
            {
                "label": day.strftime("%d/%m"),
                "value": value,
                "height": round((value / daily_max) * 100) if daily_max else 0,
            }
            for day, value in zip(daily_days, daily_values)
        ],
        "weekly": [
            {
                "label": start.strftime("%d/%m"),
                "value": value,
                "height": round((value / weekly_max) * 100) if weekly_max else 0,
            }
            for start, value in zip(week_starts, weekly_values)
        ],
        "daily_max": daily_max,
        "weekly_max": weekly_max,
    }


def equipment_rows():
    rows = []
    for row in sheet_rows("Inventario Equipos"):
        padded = (row + [""] * 8)[:8]
        rows.append(
            {
                "id": clean(padded[0]),
                "imei": clean(padded[1]),
                "modelo": clean(padded[2]),
                "gama": clean(padded[3]),
                "estado_equipo": clean(padded[4]),
                "situacion": clean(padded[5]),
                "fecha_alta": fmt_date(padded[6]),
                "observaciones": clean(padded[7]),
            }
        )
    return rows


def equipment_by_imei(imei: str):
    for row in equipment_rows():
        if row["imei"] == clean(imei):
            return row
    return None


def visible_assignments(query: str = ""):
    wb = book()
    q = query.casefold().strip()
    inv_by_id = {}
    if "Inventario" in wb.sheetnames:
        for row in wb["Inventario"].iter_rows(min_row=2):
            key = clean(row[0].value)
            if key:
                inv_by_id[key] = row
    equip_by_imei = {}
    if "Inventario Equipos" in wb.sheetnames:
        for row in wb["Inventario Equipos"].iter_rows(min_row=2):
            key = clean(row[1].value)
            if key:
                equip_by_imei[key] = row

    rows = []
    for ar in wb["Asignaciones"].iter_rows(min_row=2):
        if clean(ar[13].value).casefold() != "activo":
            continue
        eqid = clean(ar[1].value)
        inv = inv_by_id.get(eqid)
        imei = clean(ar[7].value or (inv[12].value if inv else ""))
        equip = equip_by_imei.get(imei)
        visible = {
            "operador": clean(ar[10].value or (inv[1].value if inv else "")),
            "razon_social": clean(inv[2].value if inv else ""),
            "nombre": clean(ar[2].value or (inv[3].value if inv else "")),
            "correo": clean(ar[4].value or (inv[4].value if inv else "")),
            "cargo": clean(ar[5].value or (inv[5].value if inv else "")),
            "dni": clean(ar[3].value or (inv[6].value if inv else "")),
            "celular": clean(ar[9].value or (inv[7].value if inv else "")),
            "plan": clean(inv[8].value if inv else ""),
            "tipo_plan": clean(inv[9].value if inv else ""),
            "modelo": clean(ar[6].value or (inv[10].value if inv else "")),
            "gama": clean(inv[11].value if inv else ""),
            "imei": imei,
            "situacion": clean(equip[5].value if equip else (inv[13].value if inv else "")),
        }
        haystack = " ".join(visible.values()).casefold()
        if not q or q in haystack:
            rows.append(visible)
    return rows


def available_imeis():
    return [row for row in equipment_rows() if row["situacion"] == "Disponible"]


def movement_rows():
    labels = {
        "CARGA MASIVA": "Carga Masiva",
        "ASIGNACION": "Asignacion",
        "MANTENIMIENTO": "Mantenimiento",
        "REHABILITACION": "Rehabilitacion",
        "BAJA": "Baja",
        "ELIMINACION": "Eliminacion",
        "RECUPERACION": "Recuperacion",
    }
    rows = []
    for row in sheet_rows("Movimientos"):
        padded = (row + [""] * 7)[:7]
        raw_type = clean(padded[2])
        detail = clean(padded[6])
        if raw_type in ("ASIGNACION", "ELIMINACION"):
            detail = re.sub(r"\s*\|?\s*IMEI\s*[:=-]?\s*\d{14,20}\b", "", detail, flags=re.IGNORECASE)
            detail = re.sub(r"^\s*\d{14,20}\s+", "", detail).replace("  ", " ").strip(" |")
        rows.append(
            {
                "tipo": labels.get(raw_type, raw_type),
                "fecha": fmt_date(padded[3]),
                "usuario": clean(padded[4]),
                "imei": clean(padded[5]),
                "detalle": detail,
            }
        )
    return rows


def acta_rows():
    rows = []
    for row in sheet_rows("Actas de Entrega"):
        padded = (row + [""] * len(ACTAS_ENTREGA_HEADERS))[: len(ACTAS_ENTREGA_HEADERS)]
        rows.append(dict(zip(ACTAS_ENTREGA_HEADERS, [clean(x) for x in padded])))
    return rows


def user_rows():
    rows = []
    for row in sheet_rows("Usuarios"):
        padded = (row + [""] * 3)[:3]
        rows.append({"usuario": clean(padded[0]), "rol": clean(padded[1])})
    return rows


def flash_form_error(message: str, endpoint: str = "dashboard"):
    flash(message, "error")
    return redirect(url_for(endpoint))


def create_equipment(data):
    imei = clean(data.get("imei"))
    if not imei.isdigit() or len(imei) not in (14, 15):
        raise ValueError("El IMEI debe tener 14 o 15 digitos.")
    wb = book()
    ws = wb["Inventario Equipos"]
    if any(clean(row[1].value) == imei for row in ws.iter_rows(min_row=2)):
        raise ValueError("El IMEI ya existe.")
    next_id = max([int_or_zero(row[0].value) for row in ws.iter_rows(min_row=2)], default=0) + 1
    ws.append(
        [
            next_id,
            imei,
            clean(data.get("modelo")),
            clean(data.get("gama")) or GAMAS[0],
            clean(data.get("estado_equipo")) or EQUIP_STATES[0],
            "Disponible",
            now(),
            clean(data.get("observaciones")),
        ]
    )
    wb.save(EXCEL)
    log(session["user"], next_id, "ALTA EQUIPO", "IMEI " + imei, imei)


def assign_equipment(data):
    imei = clean(data.get("imei"))
    nombre = clean(data.get("nombre"))
    dni = clean(data.get("dni"))
    if not imei or not nombre or not dni:
        raise ValueError("Seleccione IMEI e indique Nombre y DNI.")
    equipo = equipment_by_imei(imei)
    if not equipo or equipo["situacion"].casefold() != "disponible":
        raise ValueError("El IMEI no esta disponible.")

    wb = book()
    for row in wb["Inventario Equipos"].iter_rows(min_row=2):
        if clean(row[1].value) == imei:
            row[5].value = "Asignado"
            break

    ws = wb["Inventario"]
    eqid = max([int_or_zero(row[0].value) for row in ws.iter_rows(min_row=2)], default=0) + 1
    inv_row = [
        eqid,
        clean(data.get("operador")) or OPERADORES[0],
        clean(data.get("razon_social")) or RAZONES[0],
        nombre,
        clean(data.get("correo")),
        clean(data.get("cargo")),
        dni,
        clean(data.get("celular")),
        clean(data.get("plan")) or PLANES[0],
        clean(data.get("tipo_plan")) or TIPOS_PLAN[0],
        equipo["modelo"],
        equipo["gama"],
        imei,
        "Asignado",
        equipo["estado_equipo"],
        now(),
        "",
        "",
    ]
    ws.append(inv_row)

    wa = wb["Asignaciones"]
    aid = max([int_or_zero(row[0].value) for row in wa.iter_rows(min_row=2)], default=0) + 1
    wa.append(
        [
            aid,
            eqid,
            inv_row[3],
            inv_row[6],
            inv_row[4],
            inv_row[5],
            inv_row[10],
            inv_row[12],
            inv_row[14],
            inv_row[7],
            inv_row[1],
            now(),
            "",
            "Activo",
            "Asignacion desde Inventario Equipos / IMEI",
        ]
    )
    wb.save(EXCEL)
    log(session["user"], eqid, "ASIGNACION", "Asignado a " + nombre, imei)


def find_inventory_by_imei(wb, imei: str):
    for row in wb["Inventario"].iter_rows(min_row=2):
        if clean(row[12].value) == clean(imei):
            return row
    return None


def change_state(imei: str, state: str, observation: str = "", release: bool = False):
    wb = book()
    inv = find_inventory_by_imei(wb, imei)
    if inv is None:
        raise ValueError("No se encontro el IMEI en asignaciones.")
    current_state = clean(inv[13].value)
    if current_state.casefold() == "baja" and state.casefold() != "baja":
        raise ValueError("El equipo esta dado de Baja y no puede cambiar de estado.")

    eqid = clean(inv[0].value)
    mark_as_seminew = release or state.casefold() == "mantenimiento"
    for row in wb["Inventario"].iter_rows(min_row=2):
        if clean(row[12].value) == imei:
            row[13].value = state
            if mark_as_seminew:
                row[14].value = "Equipo seminuevo"
            row[16].value = now()
            if observation:
                row[17].value = observation
    for row in wb["Inventario Equipos"].iter_rows(min_row=2):
        if clean(row[1].value) == imei:
            row[5].value = state
            if mark_as_seminew:
                row[4].value = "Equipo seminuevo"
            if observation:
                row[7].value = observation

    wa = wb["Asignaciones"]
    deleted = 0
    for idx in range(wa.max_row, 1, -1):
        row = wa[idx]
        row_eqid = clean(row[1].value)
        row_imei = clean(row[7].value)
        row_status = clean(row[13].value).casefold()
        if row_status == "activo" and (row_eqid == eqid or row_imei == imei):
            wa.delete_rows(idx, 1)
            deleted += 1
    if deleted == 0:
        for idx in range(wa.max_row, 1, -1):
            row = wa[idx]
            if clean(row[1].value) == eqid or clean(row[7].value) == imei:
                wa.delete_rows(idx, 1)
                deleted += 1

    wb.save(EXCEL)
    action = "RECUPERACION" if release else ("MANTENIMIENTO" if state == "Mantenimiento" else "BAJA" if state == "Baja" else "CAMBIO ESTADO")
    detail = state + (" | " + observation if observation else "")
    log(session["user"], int_or_zero(eqid), action, detail, imei)


def delete_assignment(imei: str):
    wb = book()
    inv = find_inventory_by_imei(wb, imei)
    eqid = clean(inv[0].value) if inv is not None else ""
    wa = wb["Asignaciones"]
    deleted = 0
    for idx in range(wa.max_row, 1, -1):
        row = wa[idx]
        if clean(row[7].value) == imei or (eqid and clean(row[1].value) == eqid):
            wa.delete_rows(idx, 1)
            deleted += 1
    if deleted == 0:
        raise ValueError("No se encontro el registro seleccionado en Asignaciones.")
    for row in wb["Inventario"].iter_rows(min_row=2):
        if clean(row[12].value) == imei:
            row[13].value = "Disponible"
            row[16].value = now()
    for row in wb["Inventario Equipos"].iter_rows(min_row=2):
        if clean(row[1].value) == imei:
            row[5].value = "Disponible"
            row[7].value = "Eliminacion manual de asignacion"
    wb.save(EXCEL)
    log(session["user"], int_or_zero(eqid), "ELIMINACION", "Asignacion eliminada", imei)


def rehabilitate_equipment(imei: str, observation: str):
    if not observation:
        raise ValueError("Debe ingresar una observacion para rehabilitar el equipo.")
    equipo = equipment_by_imei(imei)
    if not equipo:
        raise ValueError("No se encontro el equipo.")
    if equipo["situacion"] == "Baja":
        raise ValueError("Un equipo dado de Baja no puede ser rehabilitado.")
    if equipo["situacion"] != "Mantenimiento":
        raise ValueError("Solo se pueden rehabilitar equipos en Mantenimiento.")
    wb = book()
    for row in wb["Inventario Equipos"].iter_rows(min_row=2):
        if clean(row[1].value) == imei:
            row[5].value = "Disponible"
            row[7].value = observation
            break
    wb.save(EXCEL)
    log(session["user"], imei, "REHABILITACION", "Equipo rehabilitado: " + observation, imei)


def import_equipment(path: Path):
    try:
        src = load_workbook(path, data_only=True)
        sw = src.active
        headers = [clean(cell.value) for cell in sw[1]]
    except Exception as exc:
        raise ValueError(f"No se pudo leer el Excel: {exc}") from exc
    if "IMEI" not in headers:
        raise ValueError("El Excel debe contener la columna IMEI.")
    if "Marca y Modelo" not in headers and "Modelo Asignado" not in headers and not ("Marca" in headers and "Modelo" in headers):
        raise ValueError("El Excel debe contener IMEI y Marca y Modelo.")

    wb = book()
    ws = wb["Inventario Equipos"]
    existing = {clean(row[1].value) for row in ws.iter_rows(min_row=2) if row[1].value}
    next_id = max([int_or_zero(row[0].value) for row in ws.iter_rows(min_row=2)], default=0) + 1
    count = 0
    skipped = 0
    for values in sw.iter_rows(min_row=2, values_only=True):
        data = dict(zip(headers, values))
        imei = clean(data.get("IMEI"))
        if imei.endswith(".0"):
            imei = imei[:-2]
        if not imei.isdigit() or len(imei) not in (14, 15) or imei in existing:
            skipped += 1
            continue
        combined = clean(data.get("Marca y Modelo") or data.get("Modelo Asignado"))
        if not combined:
            combined = " ".join(part for part in (clean(data.get("Marca")), clean(data.get("Modelo"))) if part)
        ws.append(
            [
                next_id,
                imei,
                combined,
                clean(data.get("Gama Celular")) or "NA",
                clean(data.get("Estado de equipo")) or "Equipo nuevo",
                "Disponible",
                now(),
                "Carga masiva",
            ]
        )
        existing.add(imei)
        count += 1
        next_id += 1
    wb.save(EXCEL)
    log(session["user"], "", "CARGA MASIVA", f"Equipos importados: {count}; omitidos: {skipped}")
    return count, skipped


def export_report_path() -> Path:
    src = book()
    out = Workbook()
    out.remove(out.active)
    hidden_by_sheet = {
        "Inventario": {"id"},
        "Inventario Equipos": {"id equipo"},
        "Asignaciones": {"id asignacion", "id asignación", "id equipo"},
        "Movimientos": {"id movimiento", "id equipo"},
        "Actas": {"id equipo"},
    }
    allowed = ["Inventario", "Inventario Equipos", "Asignaciones"]
    if session.get("role") == "Administrador":
        allowed += ["Movimientos", "Usuarios", "Actas"]
    for name in allowed:
        if name not in src.sheetnames:
            continue
        sw = src[name]
        dw = out.create_sheet(name)
        headers = [clean(cell.value) for cell in sw[1]]
        omit = hidden_by_sheet.get(name, set())
        keep = [idx for idx, h in enumerate(headers) if h.casefold() not in omit]
        for out_col, src_idx in enumerate(keep, start=1):
            dw.cell(1, out_col, sw.cell(1, src_idx + 1).value)
        for src_row in range(2, sw.max_row + 1):
            for out_col, src_idx in enumerate(keep, start=1):
                src_cell = sw.cell(src_row, src_idx + 1)
                dst = dw.cell(src_row, out_col, src_cell.value)
                dst.number_format = src_cell.number_format
        dw.freeze_panes = "A2"
        dw.auto_filter.ref = dw.dimensions
    fd, name = tempfile.mkstemp(prefix="Gestion_Celulares_Exportacion_", suffix=".xlsx")
    os.close(fd)
    path = Path(name)
    out.save(path)
    return path


def generate_acta_pdf_reportlab(pdf: Path, data: dict):
    pdf.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    normal = styles["Normal"]
    normal.fontName = "Helvetica"
    normal.fontSize = 10
    normal.leading = 14
    title = ParagraphStyle(
        "ActaTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        alignment=1,
        textColor=colors.HexColor("#1F4E78"),
        spaceAfter=12,
    )
    section = ParagraphStyle(
        "Section",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        textColor=colors.HexColor("#1F4E78"),
        spaceBefore=8,
        spaceAfter=6,
    )

    doc = SimpleDocTemplate(
        str(pdf),
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title="Acta de Entrega",
    )

    story = [
        Paragraph("ACTA DE ENTREGA DE EQUIPO CELULAR", title),
        Paragraph(
            "Por medio de la presente se deja constancia de la entrega del equipo celular "
            "y linea corporativa al colaborador indicado, para uso exclusivo de sus funciones.",
            normal,
        ),
        Spacer(1, 10),
        Paragraph("Datos del colaborador", section),
    ]

    employee_rows = [
        ["Nombre y Apellidos", data["nombre"]],
        ["Cargo", data["cargo"]],
        ["DNI", data["dni"]],
        ["Fecha de entrega", data["fecha_entrega"]],
    ]
    equipment_rows_pdf = [
        ["Marca y Modelo", data["modelo"]],
        ["IMEI", data["imei"]],
        ["Estado de equipo", data["estado_equipo"]],
        ["Numero celular", data["celular"]],
        ["Operador", data["operador"]],
    ]

    def info_table(rows):
        table = Table(rows, colWidths=[5.2 * cm, 11.8 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9E2F3")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF1F7")),
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        return table

    story.append(info_table(employee_rows))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Datos del equipo", section))
    story.append(info_table(equipment_rows_pdf))
    story.append(Spacer(1, 18))
    story.append(
        Paragraph(
            "El colaborador declara recibir el equipo descrito en buen estado y se compromete "
            "a conservarlo, usarlo responsablemente y devolverlo cuando la empresa lo solicite.",
            normal,
        )
    )
    story.append(Spacer(1, 34))
    signatures = Table(
        [
            ["____________________________", "____________________________"],
            ["Entrega", "Recibe"],
            ["", data["nombre"]],
        ],
        colWidths=[8 * cm, 8 * cm],
    )
    signatures.setStyle(
        TableStyle(
            [
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#1F4E78")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(signatures)
    doc.build(story)


def generate_acta_pdf(imei: str, user: str) -> Path:
    wb = book()
    wa = wb["Asignaciones"]
    ca = header_map(wa)
    modelo_col = "marca y modelo" if "marca y modelo" in ca else ("modelo asignado" if "modelo asignado" in ca else None)
    # Accept either accented or unaccented versions for migrated workbooks.
    has_phone = "numero celular" in ca or "número celular" in ca
    missing = [h for h in ["imei", "nombre y apellido", "cargo", "dni", "operador", "estado", "id equipo"] if h not in ca]
    if not has_phone:
        missing.append("numero celular")
    if modelo_col is None:
        missing.append("marca y modelo")
    if missing:
        raise ValueError("Faltan columnas requeridas en Asignaciones: " + ", ".join(missing))

    phone_key = "número celular" if "número celular" in ca else "numero celular"
    asign = None
    for row in wa.iter_rows(min_row=2):
        if clean(row[ca["imei"]].value) == imei and clean(row[ca["estado"]].value).casefold() == "activo":
            asign = row
            break
    if asign is None:
        raise ValueError("No se encontro una asignacion activa para el IMEI seleccionado.")

    estado_equipo = ""
    if "Inventario Equipos" in wb.sheetnames:
        we = wb["Inventario Equipos"]
        ce = header_map(we)
        if "imei" in ce and "estado de equipo" in ce:
            for row in we.iter_rows(min_row=2):
                if clean(row[ce["imei"]].value) == imei:
                    estado_equipo = clean(row[ce["estado de equipo"]].value)
                    break

    inv_match = None
    if "Inventario" in wb.sheetnames:
        wi = wb["Inventario"]
        ci = header_map(wi)
        eqid = clean(asign[ca["id equipo"]].value)
        for row in wi.iter_rows(min_row=2):
            if ("imei" in ci and clean(row[ci["imei"]].value) == imei) or ("id" in ci and eqid and clean(row[ci["id"]].value) == eqid):
                if inv_match is None or clean(row[ci.get("estado", 0)].value).casefold() == "asignado":
                    inv_match = row
        if not estado_equipo and inv_match is not None and "estado de equipo" in ci:
            estado_equipo = clean(inv_match[ci["estado de equipo"]].value)

    nombre = clean(asign[ca["nombre y apellido"]].value)
    cargo = clean(asign[ca["cargo"]].value)
    dni = clean(asign[ca["dni"]].value)
    modelo = clean(asign[ca[modelo_col]].value)
    imei_acta = clean(asign[ca["imei"]].value)
    celular = clean(asign[ca[phone_key]].value)
    operador = clean(asign[ca["operador"]].value)
    if inv_match is not None:
        ci = header_map(wb["Inventario"])
        if "numero celular" in ci and not celular:
            celular = clean(inv_match[ci["numero celular"]].value)
        if "número celular" in ci and not celular:
            celular = clean(inv_match[ci["número celular"]].value)
        if "modelo asignado" in ci and not modelo:
            modelo = clean(inv_match[ci["modelo asignado"]].value)

    if not nombre or not dni or not imei_acta:
        raise ValueError("La asignacion debe contener Nombre, DNI e IMEI.")
    if not celular:
        raise ValueError("La asignacion seleccionada no tiene numero celular.")
    if not estado_equipo:
        raise ValueError("No se encontro el Estado de equipo para el IMEI seleccionado.")
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"No se encontro la plantilla: {TEMPLATE.name}")

    n = max([int_or_zero(row[0].value) for row in wb["Actas"].iter_rows(min_row=2)], default=0) + 1
    safe_name = "".join(c for c in nombre if c not in '<>:/\\|?*').strip() or "Sin Nombre"
    pdf = ACTAS / f"ACTA DE ENTREGA - {safe_name}.pdf"
    fecha_entrega = datetime.now().strftime("%d/%m/%Y")
    repl = {
        "{{Nombre y Apellidos}}": nombre,
        "{{Cargo}}": cargo,
        "{{DNI}}": dni,
        "{{Modelo Asignado}}": modelo,
        "{{Numero de IMEI}}": imei_acta,
        "{{Número de IMEI}}": imei_acta,
        "{{Estado de equipo}}": estado_equipo,
        "{{Numero celular}}": celular,
        "{{Número celular}}": celular,
        "{{Operador}}": operador,
        "{{Fecha de entrega}}": fecha_entrega,
    }

    generate_acta_pdf_reportlab(
        pdf,
        {
            "nombre": nombre,
            "cargo": cargo,
            "dni": dni,
            "modelo": modelo,
            "imei": imei_acta,
            "estado_equipo": estado_equipo,
            "celular": celular,
            "operador": operador,
            "fecha_entrega": fecha_entrega,
        },
    )

    wb = book()
    fecha_generacion = datetime.now().strftime("%d/%m/%Y")
    wb["Actas"].append([n, clean(asign[ca["id equipo"]].value), "Entrega", nombre, dni, fecha_generacion, pdf.name])
    if "Actas de Entrega" not in wb.sheetnames:
        ws = wb.create_sheet("Actas de Entrega")
        ws.append(ACTAS_ENTREGA_HEADERS)
    wb["Actas de Entrega"].append([fecha_generacion, user, nombre, dni, fecha_entrega, pdf.name])
    wb.save(EXCEL)
    save_pdf_backup(pdf.name, pdf)
    return pdf

    # Flujo historico local con Microsoft Word. Se conserva como referencia, pero
    # la version web usa ReportLab para funcionar tambien en Vercel.
    temp_ps = ACTAS / f".generar_acta_word_{n:05d}.ps1"
    payload = ACTAS / f".acta_{n:05d}.json"
    temp_docx = ACTAS / f".acta_template_{n:05d}.docx"
    try:
        if platform.system() != "Windows":
            raise RuntimeError("La generacion del acta requiere Windows con Microsoft Word. En Vercel esta accion queda deshabilitada para pruebas.")
        from app_celulares import materialize_acta_resources

        materialize_acta_resources(temp_ps)
        payload.write_text(json.dumps(repl, ensure_ascii=False), encoding="utf-8")
        with zipfile.ZipFile(TEMPLATE, "r") as zin, zipfile.ZipFile(temp_docx, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                    text = data.decode("utf-8", errors="ignore")
                    for marker, value in repl.items():
                        text = text.replace(marker, str(value or ""))
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(temp_ps),
                "-Template",
                str(temp_docx),
                "-OutputPdf",
                str(pdf),
                "-DataJson",
                str(payload),
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        if not pdf.exists():
            raise RuntimeError("Microsoft Word no genero el PDF esperado.")
    finally:
        for path in (payload, temp_ps, temp_docx):
            try:
                path.unlink()
            except Exception:
                pass

    wb = book()
    fecha_generacion = datetime.now().strftime("%d/%m/%Y")
    wb["Actas"].append([n, clean(asign[ca["id equipo"]].value), "Entrega", nombre, dni, fecha_generacion, pdf.name])
    if "Actas de Entrega" not in wb.sheetnames:
        ws = wb.create_sheet("Actas de Entrega")
        ws.append(ACTAS_ENTREGA_HEADERS)
    wb["Actas de Entrega"].append([fecha_generacion, user, nombre, dni, fecha_entrega, pdf.name])
    wb.save(EXCEL)
    return pdf


@app.context_processor
def inject_globals():
    return {
        "current_user": current_user(),
        "level": level(),
        "roles": ROLES,
        "states": STATES,
        "operadores": OPERADORES,
        "razones": RAZONES,
        "planes": PLANES,
        "tipos_plan": TIPOS_PLAN,
        "gamas": GAMAS,
        "equip_states": EQUIP_STATES,
    }


@app.route("/", methods=["GET"])
def index():
    if current_user():
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = clean(request.form.get("usuario"))
        password = request.form.get("clave", "")
        for row in sheet_rows("Usuarios"):
            padded = (row + [""] * 3)[:3]
            if clean(padded[0]) == username and clean(padded[2]) == sha(password):
                session.clear()
                session["user"] = clean(padded[0])
                session["role"] = clean(padded[1])
                return redirect(url_for("dashboard"))
        flash("Usuario o clave incorrectos.", "error")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
def dashboard():
    return render_template(
        "dashboard.html",
        active_view="dashboard",
        counts=dashboard_counts(),
        charts=movement_chart_data(),
    )


@app.route("/asignaciones")
@login_required
def assignments_view():
    q = clean(request.args.get("q"))
    return render_template(
        "asignaciones.html",
        active_view="asignaciones",
        assignments=visible_assignments(q),
        available=available_imeis(),
        q=q,
    )


@app.route("/equipos")
@login_required
def equipment_view():
    return render_template(
        "equipos.html",
        active_view="equipos",
        equipment=equipment_rows(),
    )


@app.route("/historial")
@login_required
def history_view():
    return render_template("historial.html", active_view="historial", movements=movement_rows())


@app.route("/actas")
@login_required
def actas_view():
    return render_template("actas.html", active_view="actas", actas=acta_rows())


@app.route("/usuarios")
@login_required
@role_required(3)
def users_view():
    return render_template("usuarios.html", active_view="usuarios", users=user_rows())


@app.route("/equipment", methods=["POST"])
@login_required
@role_required(2)
def equipment_create():
    try:
        create_equipment(request.form)
        flash("Equipo registrado correctamente.", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("equipment_view"))


@app.route("/equipment/import", methods=["POST"])
@login_required
@role_required(2)
def equipment_import():
    upload = request.files.get("archivo")
    if not upload or not upload.filename:
        flash("Seleccione un archivo Excel.", "error")
        return redirect(url_for("equipment_view"))
    suffix = Path(upload.filename).suffix.lower()
    if suffix != ".xlsx":
        flash("El archivo debe ser .xlsx.", "error")
        return redirect(url_for("equipment_view"))
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / secure_filename(upload.filename)
        upload.save(path)
        try:
            count, skipped = import_equipment(path)
            flash(f"Se importaron {count} equipos. Registros omitidos: {skipped}.", "success")
        except Exception as exc:
            flash(str(exc), "error")
    return redirect(url_for("equipment_view"))


@app.route("/equipment/<imei>/rehabilitate", methods=["POST"])
@login_required
@role_required(2)
def equipment_rehabilitate(imei):
    try:
        rehabilitate_equipment(imei, clean(request.form.get("observacion")))
        flash("Equipo rehabilitado correctamente.", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("equipment_view"))


@app.route("/assignments", methods=["POST"])
@login_required
@role_required(2)
def assignment_create():
    try:
        assign_equipment(request.form)
        flash("Asignacion creada correctamente.", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("assignments_view"))


@app.route("/assignments/<imei>/state", methods=["POST"])
@login_required
@role_required(2)
def assignment_state(imei):
    action = clean(request.form.get("action"))
    observation = clean(request.form.get("observacion"))
    try:
        if action == "recuperar":
            if not observation:
                raise ValueError("Debe ingresar una observacion para recuperar el equipo.")
            change_state(imei, "Disponible", observation, release=True)
        elif action == "mantenimiento":
            if not observation:
                raise ValueError("Debe ingresar una observacion para mantenimiento.")
            change_state(imei, "Mantenimiento", observation)
        elif action == "baja":
            reason = observation or clean(request.form.get("motivo")) or "Robo"
            change_state(imei, "Baja", reason)
        else:
            raise ValueError("Accion no reconocida.")
        flash("Estado actualizado correctamente.", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("assignments_view"))


@app.route("/assignments/<imei>/delete", methods=["POST"])
@login_required
@role_required(3)
def assignment_delete(imei):
    try:
        delete_assignment(imei)
        flash("Asignacion eliminada correctamente.", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("assignments_view"))


@app.route("/assignments/<imei>/acta", methods=["POST"])
@login_required
@role_required(2)
def assignment_acta(imei):
    try:
        pdf = generate_acta_pdf(imei, session["user"])
        flash(f"Acta generada: {pdf.name}", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("assignments_view"))


@app.route("/actas/<path:filename>")
@login_required
def acta_download(filename):
    safe = Path(filename).name
    path = ACTAS / safe
    if not path.exists():
        restored = restore_pdf_from_backup(safe)
        if restored and restored.exists():
            path = restored
    if not path.exists():
        flash("No se encontro el PDF solicitado.", "error")
        return redirect(url_for("actas_view"))
    return send_file(path, as_attachment=False)


@app.route("/export")
@login_required
@role_required(2)
def export():
    path = export_report_path()
    return send_file(path, as_attachment=True, download_name="Gestion_Celulares_Exportacion.xlsx")


@app.route("/users", methods=["POST"])
@login_required
@role_required(3)
def user_create():
    username = clean(request.form.get("usuario"))
    password = request.form.get("clave", "")
    role = clean(request.form.get("rol")) or "Consulta"
    if not username or not password or role not in ROLES:
        flash("Complete usuario, clave y rol.", "error")
        return redirect(url_for("users_view"))
    wb = book()
    ws = wb["Usuarios"]
    if any(clean(row[0].value) == username for row in ws.iter_rows(min_row=2)):
        flash("El usuario ya existe.", "error")
        return redirect(url_for("users_view"))
    ws.append([username, role, sha(password)])
    wb.save(EXCEL)
    log(session["user"], "", "USUARIO", "Creado " + username)
    flash("Usuario creado correctamente.", "success")
    return redirect(url_for("users_view"))


@app.route("/users/<username>/delete", methods=["POST"])
@login_required
@role_required(3)
def user_delete(username):
    if username in (session.get("user"), "admin"):
        flash("No puede eliminarse ese usuario.", "error")
        return redirect(url_for("users_view"))
    wb = book()
    ws = wb["Usuarios"]
    for idx in range(2, ws.max_row + 1):
        if clean(ws.cell(idx, 1).value) == username:
            ws.delete_rows(idx)
            break
    wb.save(EXCEL)
    log(session["user"], "", "USUARIO", "Eliminado " + username)
    flash("Usuario eliminado correctamente.", "success")
    return redirect(url_for("users_view"))


if __name__ == "__main__":
    prepare_database()
    app.run(host="127.0.0.1", port=5000, debug=os.environ.get("SIGE_DEBUG") == "1")
